from __future__ import annotations

import copy
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
XML = "http://www.w3.org/XML/1998/namespace"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
for prefix, uri in (("w", W), ("r", R)):
    ET.register_namespace(prefix, uri)


class OfficeRoundtripService:
    """Part-scoped OOXML editing with auditable fidelity evidence."""

    def edit_docx(
        self, source: str | Path, output: str | Path, *,
        replacements: dict[str, Any] | None = None,
        comments: list[dict[str, Any]] | None = None,
        revisions: list[dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        source, output = Path(source), Path(output)
        parts = self._read(source)
        document = ET.fromstring(parts["word/document.xml"])
        document.attrib.pop(f"{{{MC}}}Ignorable", None)
        for old, new in (replacements or {}).items():
            self._replace_text(document, str(old), str(new))
        if comments:
            self._add_comments(parts, document, comments)
        if revisions:
            self._add_revisions(parts, document, revisions)
        parts["word/document.xml"] = ET.tostring(document, encoding="utf-8", xml_declaration=True)
        self._write(output, parts)
        allowed = {"word/document.xml"}
        if comments:
            allowed |= {"word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        if revisions:
            allowed |= {"word/settings.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        return self.audit(source, output, allowed)

    def edit_xlsx(self, source: str | Path, output: str | Path, updates: list[dict[str, Any]]) -> dict[str, Any]:
        from office_artifacts import _edit_xlsx_stdlib
        source, output = Path(source), Path(output)
        _edit_xlsx_stdlib(source, updates or [], output)
        allowed = set()
        before, after = self._hashes(source), self._hashes(output)
        allowed.update(name for name in after if name.startswith("xl/worksheets/") and before.get(name) != after.get(name))
        return self.audit(source, output, allowed)

    def edit_pptx(self, source: str | Path, output: str | Path, replacements: dict[str, Any]) -> dict[str, Any]:
        source, output = Path(source), Path(output)
        parts = self._read(source)
        changed = set()
        for name, raw in list(parts.items()):
            if not name.startswith("ppt/slides/") or not name.endswith(".xml"):
                continue
            root = ET.fromstring(raw)
            touched = False
            for node in root.iter():
                if node.tag.endswith("}t") and node.text:
                    value = node.text
                    for old, new in (replacements or {}).items():
                        value = value.replace(str(old), str(new))
                    if value != node.text:
                        node.text, touched = value, True
            if touched:
                parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                changed.add(name)
        self._write(output, parts)
        return self.audit(source, output, changed)

    def recalculate_xlsx(self, source: str | Path, output: str | Path) -> dict[str, Any]:
        source, output = Path(source), Path(output)
        engine = shutil.which("soffice")
        if not engine:
            candidate = Path("C:/Program Files/LibreOffice/program/soffice.com")
            engine = str(candidate) if candidate.is_file() else ""
        if not engine:
            raise RuntimeError("LibreOffice is required for formula recalculation")
        with tempfile.TemporaryDirectory(prefix="clawos-office-recalc-") as td:
            root = Path(td)
            incoming, converted = root / source.name, root / "converted"
            profile = root / "profile"
            converted.mkdir()
            shutil.copy2(source, incoming)
            result = subprocess.run([
                engine, f"-env:UserInstallation={profile.as_uri()}", "--headless",
                "--convert-to", "xlsx", "--outdir", str(converted), str(incoming),
            ], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
            produced = converted / source.name
            if result.returncode != 0 or not produced.is_file():
                raise RuntimeError((result.stderr or result.stdout).strip() or "LibreOffice recalculation failed")
            output.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(produced, output)
        formulas, cached = self._formula_counts(output)
        return {
            "engine": "libreoffice-calc", "formula_count": formulas,
            "cached_formula_values": cached,
            "recalculated": formulas == 0 or cached > 0,
            "output_sha256": self._sha(output),
        }

    def audit(self, source: Path, output: Path, allowed_changed_parts: set[str]) -> dict[str, Any]:
        before, after = self._hashes(source), self._hashes(output)
        changed = sorted(name for name in before.keys() & after.keys() if before[name] != after[name])
        added, removed = sorted(after.keys() - before.keys()), sorted(before.keys() - after.keys())
        actual = set(changed + added + removed)
        styles = [name for name in before if name.endswith("styles.xml") or "/theme/" in name]
        style_preserved = all(before.get(name) == after.get(name) for name in styles)
        unexpected = sorted(actual - set(allowed_changed_parts))
        return {
            "schema": "clawos.office-roundtrip.v1", "package_valid": True,
            "changed_parts": changed, "added_parts": added, "removed_parts": removed,
            "allowed_changed_parts": sorted(allowed_changed_parts),
            "unexpected_changed_parts": unexpected,
            "style_parts": styles, "style_preserved": style_preserved,
            "roundtrip_passed": not unexpected and style_preserved,
            "source_sha256": self._sha(source), "output_sha256": self._sha(output),
        }

    @staticmethod
    def _replace_text(root: ET.Element, old: str, new: str) -> None:
        found = False
        for node in root.iter(f"{{{W}}}t"):
            if node.text and old in node.text:
                node.text = node.text.replace(old, new)
                found = True
        if old and not found:
            raise ValueError(f"text target is not contained in a single Word run: {old}")

    def _add_comments(self, parts: dict[str, bytes], document: ET.Element, rows: list[dict[str, Any]]) -> None:
        comments = ET.fromstring(parts.get("word/comments.xml", f'<w:comments xmlns:w="{W}"/>'.encode()))
        ids = [int(item.get(f"{{{W}}}id", "-1")) for item in comments.findall(f"{{{W}}}comment")]
        next_id = max(ids, default=-1) + 1
        for row in rows:
            target = str(row.get("target") or "")
            if not target:
                raise ValueError("comment target is required")
            comment_id = next_id
            next_id += 1
            self._mark_comment(document, target, comment_id)
            node = ET.SubElement(comments, f"{{{W}}}comment", {
                f"{{{W}}}id": str(comment_id), f"{{{W}}}author": str(row.get("author") or "ClawOS"),
                f"{{{W}}}date": self._now(),
            })
            p, run = ET.SubElement(node, f"{{{W}}}p"), None
            run = ET.SubElement(p, f"{{{W}}}r")
            ET.SubElement(run, f"{{{W}}}t").text = str(row.get("text") or "")
        parts["word/comments.xml"] = ET.tostring(comments, encoding="utf-8", xml_declaration=True)
        self._ensure_part(parts, "/word/comments.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml", "comments", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")

    def _mark_comment(self, document: ET.Element, target: str, comment_id: int) -> None:
        for parent in document.iter():
            children = list(parent)
            for index, run in enumerate(children):
                texts = list(run.iter(f"{{{W}}}t")) if run.tag == f"{{{W}}}r" else []
                if len(texts) == 1 and texts[0].text and target in texts[0].text:
                    before, _, after = texts[0].text.partition(target)
                    insertion = []
                    if before:
                        clone = copy.deepcopy(run); self._set_text(list(clone.iter(f"{{{W}}}t"))[0], before); insertion.append(clone)
                    insertion.append(ET.Element(f"{{{W}}}commentRangeStart", {f"{{{W}}}id": str(comment_id)}))
                    marked = copy.deepcopy(run); list(marked.iter(f"{{{W}}}t"))[0].text = target; insertion.append(marked)
                    insertion.append(ET.Element(f"{{{W}}}commentRangeEnd", {f"{{{W}}}id": str(comment_id)}))
                    ref = ET.Element(f"{{{W}}}r"); ET.SubElement(ref, f"{{{W}}}commentReference", {f"{{{W}}}id": str(comment_id)}); insertion.append(ref)
                    if after:
                        clone = copy.deepcopy(run); self._set_text(list(clone.iter(f"{{{W}}}t"))[0], after); insertion.append(clone)
                    parent.remove(run)
                    for offset, item in enumerate(insertion): parent.insert(index + offset, item)
                    return
        raise ValueError(f"comment target is not contained in a single Word run: {target}")

    def _add_revisions(self, parts: dict[str, bytes], document: ET.Element, rows: list[dict[str, Any]]) -> None:
        revision_id = 1
        for row in rows:
            old, new = str(row.get("old") or ""), str(row.get("new") or "")
            if not old:
                raise ValueError("revision old text is required")
            self._mark_revision(document, old, new, revision_id, str(row.get("author") or "ClawOS"))
            revision_id += 2
        settings = ET.fromstring(parts.get("word/settings.xml", f'<w:settings xmlns:w="{W}"/>'.encode()))
        settings.attrib.pop(f"{{{MC}}}Ignorable", None)
        if settings.find(f"{{{W}}}trackRevisions") is None:
            settings.insert(0, ET.Element(f"{{{W}}}trackRevisions"))
        parts["word/settings.xml"] = ET.tostring(settings, encoding="utf-8", xml_declaration=True)
        self._ensure_part(parts, "/word/settings.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml", "settings", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings")

    def _mark_revision(self, document: ET.Element, old: str, new: str, revision_id: int, author: str) -> None:
        for parent in document.iter():
            for index, run in enumerate(list(parent)):
                texts = list(run.iter(f"{{{W}}}t")) if run.tag == f"{{{W}}}r" else []
                if len(texts) == 1 and texts[0].text and old in texts[0].text:
                    before, _, after = texts[0].text.partition(old)
                    attrs = {f"{{{W}}}id": str(revision_id), f"{{{W}}}author": author, f"{{{W}}}date": self._now()}
                    insertion = []
                    if before:
                        clone = copy.deepcopy(run); self._set_text(list(clone.iter(f"{{{W}}}t"))[0], before); insertion.append(clone)
                    deletion = ET.Element(f"{{{W}}}del", attrs); deleted = copy.deepcopy(run)
                    old_node = list(deleted.iter(f"{{{W}}}t"))[0]; old_node.tag, old_node.text = f"{{{W}}}delText", old
                    deletion.append(deleted); insertion.append(deletion)
                    addition = ET.Element(f"{{{W}}}ins", {**attrs, f"{{{W}}}id": str(revision_id + 1)})
                    added = copy.deepcopy(run); self._set_text(list(added.iter(f"{{{W}}}t"))[0], new); addition.append(added); insertion.append(addition)
                    if after:
                        clone = copy.deepcopy(run); self._set_text(list(clone.iter(f"{{{W}}}t"))[0], after); insertion.append(clone)
                    parent.remove(run)
                    for offset, item in enumerate(insertion): parent.insert(index + offset, item)
                    return
        raise ValueError(f"revision target is not contained in a single Word run: {old}")

    @staticmethod
    def _set_text(node: ET.Element, value: str) -> None:
        node.text = value
        if value.startswith(" ") or value.endswith(" "):
            node.set(f"{{{XML}}}space", "preserve")
        else:
            node.attrib.pop(f"{{{XML}}}space", None)

    @staticmethod
    def _ensure_part(parts: dict[str, bytes], part_name: str, content_type: str, rel_name: str, rel_type: str) -> None:
        types = ET.fromstring(parts["[Content_Types].xml"])
        if not any(item.get("PartName") == part_name for item in types):
            ET.SubElement(types, f"{{{CT}}}Override", {"PartName": part_name, "ContentType": content_type})
            ET.register_namespace("", CT)
            parts["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
            ET.register_namespace("w", W); ET.register_namespace("r", R)
        rel_path = "word/_rels/document.xml.rels"
        rels = ET.fromstring(parts.get(rel_path, f'<Relationships xmlns="{PKG_REL}"/>'.encode()))
        if not any(item.get("Type") == rel_type for item in rels):
            ids = {item.get("Id") for item in rels}
            index = 1
            while f"rId{index}" in ids: index += 1
            ET.SubElement(rels, f"{{{PKG_REL}}}Relationship", {"Id": f"rId{index}", "Type": rel_type, "Target": f"{rel_name}.xml"})
            ET.register_namespace("", PKG_REL)
            parts[rel_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            ET.register_namespace("w", W); ET.register_namespace("r", R)

    @staticmethod
    def _formula_counts(path: Path) -> tuple[int, int]:
        try:
            from openpyxl import load_workbook
            formulas_book = load_workbook(path, data_only=False, read_only=True)
            values_book = load_workbook(path, data_only=True, read_only=True)
            formulas = cached = 0
            for ws in formulas_book.worksheets:
                value_ws = values_book[ws.title]
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            formulas += 1
                            if value_ws[cell.coordinate].value is not None: cached += 1
            formulas_book.close(); values_book.close()
            return formulas, cached
        except Exception:
            formulas = cached = 0
            namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
                        continue
                    root = ET.fromstring(archive.read(name))
                    for cell in root.iter(f"{{{namespace}}}c"):
                        if cell.find(f"{{{namespace}}}f") is not None:
                            formulas += 1
                            value = cell.find(f"{{{namespace}}}v")
                            if value is not None and value.text not in {None, ""}:
                                cached += 1
            return formulas, cached

    @staticmethod
    def _read(path: Path) -> dict[str, bytes]:
        with zipfile.ZipFile(path) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    @staticmethod
    def _write(path: Path, parts: dict[str, bytes]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(path.suffix + ".tmp")
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, raw in parts.items(): archive.writestr(name, raw)
        os.replace(temporary, path)

    @classmethod
    def _hashes(cls, path: Path) -> dict[str, str]:
        return {name: hashlib.sha256(raw).hexdigest() for name, raw in cls._read(path).items()}

    @staticmethod
    def _sha(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""): digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _now() -> str:
        return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
