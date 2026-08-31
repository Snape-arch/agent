from __future__ import annotations

import csv
import hashlib
import ipaddress
import json
import mimetypes
import os
import re
import shutil
import smtplib
import socket
import ssl
import time
import urllib.parse
import urllib.request
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from typing import Any, Callable

from office_artifacts import create_docx, create_pdf, create_pptx, create_xlsx, edit_xlsx, replace_docx_text, replace_pptx_text
from .runtime_records import RunRecord, now_iso
from .office_visual_quality_service import OfficeVisualQualityService
from .office_roundtrip_service import OfficeRoundtripService


SUPPORTED_FORMATS = {"docx", "xlsx", "pptx", "pdf", "md", "txt", "json", "csv", "eml", "ics"}


class OfficeProductivityService:
    """Unified office artifact pipeline: create/edit/inspect/validate/package/dispatch."""

    def __init__(self, store: Any | None = None, allowed_roots: list[str | Path] | None = None, smtp_sender: Callable | None = None, calendar_sender: Callable | None = None, preview_root: str | Path | None = None):
        self.store = store
        self.allowed_roots = tuple(Path(item).expanduser().resolve() for item in (allowed_roots or [Path.cwd(), Path.home() / "Desktop", Path.home() / "Documents"]))
        self.smtp_sender = smtp_sender
        self.calendar_sender = calendar_sender
        self.visual_quality = OfficeVisualQualityService()
        self.roundtrip = OfficeRoundtripService()
        configured_preview = preview_root or os.environ.get("CLAWOS_OUTPUT_DIR", "")
        self.preview_root = (Path(configured_preview).expanduser().resolve() / "office_previews") if configured_preview else None

    def quality_audit(self, file_path: str | Path, *, template_path: str | Path = "", strict_visual: bool = False) -> dict[str, Any]:
        path = self._path(file_path)
        template = self._path(template_path) if template_path else ""
        result = self.visual_quality.audit(path, template_path=template, strict_visual=strict_visual)
        self._save_run("office_visual_quality", f"Visual QA: {path.name}", result)
        return result

    def analyze_data(self, rows: list[dict[str, Any]], output_path: str | Path, *, title: str = "Data Analysis", group_by: str = "", metrics: list[str] | None = None, confirmed: bool = False) -> dict[str, Any]:
        if confirmed is not True:
            raise PermissionError("confirmed=true is required for analysis workbook writes")
        output = self._path(output_path)
        result = self.visual_quality.analyze_and_build(rows, output, title=title, group_by=group_by, metrics=metrics)
        self._save_run("office_data_analysis", title, result)
        return result

    def create(self, output_path: str | Path, *, title: str = "", sections: list[dict[str, Any]] | None = None, data: Any = None, metadata: dict[str, Any] | None = None, confirmed: bool = False) -> dict[str, Any]:
        if confirmed is not True:
            raise PermissionError("confirmed=true is required for office artifact writes")
        path = self._path(output_path)
        fmt = path.suffix.lower().lstrip(".")
        if fmt not in SUPPORTED_FORMATS - {"eml", "ics"}:
            raise ValueError(f"unsupported office format: {fmt}")
        path.parent.mkdir(parents=True, exist_ok=True)
        sections = self._expand_semantic_tables(sections or [])
        metadata = dict(metadata or {})
        metadata["semantic_requirements"] = self._semantic_requirements(sections, metadata, fmt)
        engine = "stdlib"
        if fmt == "docx":
            create_docx(path, title, sections, metadata)
            engine = self._module_engine("docx", "python-docx", "stdlib-openxml")
            if metadata.get("comments") or metadata.get("revisions"):
                reviewed = path.with_name(f".{path.stem}.reviewed{path.suffix}")
                self.roundtrip.edit_docx(
                    path, reviewed, comments=list(metadata.get("comments") or []),
                    revisions=list(metadata.get("revisions") or []),
                )
                os.replace(reviewed, path)
                engine += "+review-ooxml"
        elif fmt == "xlsx":
            sheets = data if isinstance(data, list) else (metadata or {}).get("sheets") or [{"name": "Sheet1", "rows": []}]
            create_xlsx(path, sheets)
            engine = self._module_engine("openpyxl", "openpyxl", "stdlib-openxml")
        elif fmt == "pptx":
            slides = data if isinstance(data, list) else (metadata or {}).get("slides") or [{"title": title, "bullets": []}]
            create_pptx(path, title, slides)
            engine = self._module_engine("pptx", "python-pptx", "stdlib-openxml")
        elif fmt == "pdf":
            create_pdf(path, title, sections, metadata)
            engine = "reportlab"
        elif fmt == "json":
            path.write_text(json.dumps(data if data is not None else {"title": title, "sections": sections}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        elif fmt == "csv":
            rows = data if isinstance(data, list) else []
            self._write_csv(path, rows)
        else:
            text = str(data if data is not None else self._sections_markdown(title, sections))
            path.write_text(text, encoding="utf-8")
        return self._finalize("create", path, engine, {"title": title, "format": fmt, **metadata})

    @staticmethod
    def _expand_semantic_tables(sections: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """Keep experimental groups and continuation-row parents in generated reports."""
        expanded = []
        for section in sections:
            table = section.get("table")
            from .document_evidence_policy import table_pending
            if section.get("review_required") or isinstance(table, dict) and table_pending(table):
                raise ValueError("Document evidence requires source review before export")
            if not isinstance(table, dict) or table.get("representation") not in {"ruled_region_glyph_grid", "raster_header_grid"}:
                expanded.append(section)
                continue
            groups: dict[str, list[list[str]]] = {}
            for r, values in enumerate(table["rows"][1:], 1):
                first = next(c for c in table["cells"] if c["row"] == r and c["column"] == 0)
                group = first.get("group_context", "")
                row = list(values)
                if first.get("row_context"):
                    row[0] = first["row_context"] + " / " + row[0]
                groups.setdefault(group, []).append(row)
            for index, (group, rows) in enumerate(groups.items()):
                result = dict(section) if index == 0 else {"heading": section.get("heading", "Results")}
                paragraphs = list(result.get("paragraphs") or [])
                source_label = f"Source: {table.get('source_filename') or 'original document'}, page {table.get('page', '?')}"
                paragraphs.extend(value for value in (table.get("caption"), group, source_label) if value)
                if any(c.get("annotations") for c in table["cells"]):
                    paragraphs.append("Small-type values are preserved as annotations; their statistical meaning has not been inferred.")
                result["paragraphs"] = paragraphs
                result["table"] = {"rows": [list(table["headers"]), *rows], "header_rows": 1,
                                   "source_table_id": table["id"], "source_group": group,
                                   "source_locator": table.get("locator", "")}
                if table.get("column_widths"):
                    result["table"]["column_widths"] = list(table["column_widths"])
                expanded.append(result)
        return expanded

    def edit(self, file_path: str | Path, output_path: str | Path, *, replacements: dict[str, Any] | None = None, updates: list[dict[str, Any]] | None = None, comments: list[dict[str, Any]] | None = None, revisions: list[dict[str, Any]] | None = None, recalculate: bool = False, confirmed: bool = False) -> dict[str, Any]:
        if confirmed is not True:
            raise PermissionError("confirmed=true is required for office artifact edits")
        source, output = self._path(file_path), self._path(output_path)
        if not source.is_file() or source.suffix.lower() != output.suffix.lower():
            raise ValueError("source must exist and output must use the same format")
        fmt = source.suffix.lower().lstrip(".")
        output.parent.mkdir(parents=True, exist_ok=True)
        evidence: dict[str, Any] = {}
        if fmt == "docx":
            evidence["roundtrip"] = self.roundtrip.edit_docx(
                source, output, replacements=replacements or {},
                comments=comments or [], revisions=revisions or [],
            )
            engine = "part-scoped-openxml"
        elif fmt == "xlsx":
            evidence["roundtrip"] = self.roundtrip.edit_xlsx(source, output, updates or [])
            engine = "part-scoped-openxml"
            if recalculate:
                recalculated = output.with_name(f".{output.stem}.recalculated{output.suffix}")
                evidence["recalculation"] = self.roundtrip.recalculate_xlsx(output, recalculated)
                os.replace(recalculated, output)
                engine += "+libreoffice-calc"
        elif fmt == "pptx":
            evidence["roundtrip"] = self.roundtrip.edit_pptx(source, output, replacements or {})
            engine = "part-scoped-openxml"
        elif fmt in {"md", "txt", "json", "csv"}:
            text = source.read_text(encoding="utf-8-sig")
            for old, new in (replacements or {}).items():
                text = text.replace(str(old), str(new))
            output.write_text(text, encoding="utf-8")
            engine = "stdlib"
        else:
            raise ValueError(f"editing is not supported for {fmt}")
        return self._finalize("edit", output, engine, {"source_sha256": self._sha(source), **evidence})

    def inspect(self, file_path: str | Path) -> dict[str, Any]:
        path = self._path(file_path)
        if not path.is_file():
            raise ValueError("file_path must point to a readable file")
        fmt = path.suffix.lower().lstrip(".")
        details: dict[str, Any] = {}
        if fmt == "xlsx":
            try:
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=False, data_only=False)
                details = {"sheets": [{"name": ws.title, "rows": ws.max_row, "columns": ws.max_column} for ws in wb.worksheets], "formula_count": sum(1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("="))}
                wb.close()
            except Exception as exc:
                details = self._openxml_fallback(path, "xl/workbook.xml", exc)
        elif fmt == "docx":
            try:
                from docx import Document
                doc = Document(path)
                details = {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "text_chars": sum(len(item.text) for item in doc.paragraphs)}
            except Exception as exc:
                details = self._openxml_fallback(path, "word/document.xml", exc)
        elif fmt == "pptx":
            try:
                from pptx import Presentation
                prs = Presentation(path)
                details = {"slides": len(prs.slides), "shapes": sum(len(slide.shapes) for slide in prs.slides)}
            except Exception as exc:
                details = self._openxml_fallback(path, "ppt/presentation.xml", exc)
        elif fmt == "pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(path)
                details = {"pages": len(reader.pages), "encrypted": bool(reader.is_encrypted), "text_chars": sum(len(page.extract_text() or "") for page in reader.pages[:20])}
            except Exception as exc:
                raw = path.read_bytes()
                if raw.startswith(b"%PDF-") and b"%%EOF" in raw[-2048:]:
                    details = {"pages_estimated": max(1, raw.count(b"/Type /Page") - raw.count(b"/Type /Pages")), "fallback_validation": True, "professional_engine_unavailable": type(exc).__name__}
                else:
                    details = {"open_error": f"{type(exc).__name__}: {exc}"}
        elif fmt in {"md", "txt", "json", "csv", "eml", "ics"}:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
            details = {"lines": len(text.splitlines()), "text_chars": len(text)}
            if fmt == "json":
                json.loads(text)
            if fmt == "csv":
                rows = list(csv.reader(text.splitlines()))
                details.update({"rows": len(rows), "columns": max((len(row) for row in rows), default=0)})
        else:
            with path.open("rb") as handle:
                signature = handle.read(16).hex()
            details = {"signature": signature}
        return {"path": str(path), "format": fmt, "bytes": path.stat().st_size, "sha256": self._sha(path), "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream", "valid": "open_error" not in details, "details": details}

    def meeting_package(self, meeting: dict[str, Any], output_dir: str | Path, *, recipients: list[str] | None = None, confirmed: bool = False) -> dict[str, Any]:
        if confirmed is not True:
            raise PermissionError("confirmed=true is required for meeting package writes")
        root = self._path(output_dir)
        root.mkdir(parents=True, exist_ok=True)
        title = str(meeting.get("title") or "Meeting")
        actions = meeting.get("actions") or []
        sections = [
            {"heading": "Executive summary", "paragraphs": [meeting.get("summary") or "Meeting follow-up package."]},
            {"heading": "Decisions", "bullets": meeting.get("decisions") or []},
            {"heading": "Actions", "table": [["Owner", "Action", "Due", "Evidence"], *[[item.get("owner", ""), item.get("description", ""), item.get("due_date", ""), item.get("evidence", "")] for item in actions]]},
            {"heading": "Risks", "bullets": meeting.get("risks") or []},
        ]
        artifacts = [
            self.create(root / "meeting_minutes.docx", title=title, sections=sections, confirmed=True),
            self.create(root / "executive_brief.pdf", title=title, sections=sections, confirmed=True),
            self.create(root / "actions.xlsx", title=f"{title} actions", data=[{"name": "Actions", "rows": [["owner", "action", "due", "status"], *[[item.get("owner", ""), item.get("description", ""), item.get("due_date", ""), item.get("status", "open")] for item in actions]]}], confirmed=True),
            self.create(root / "meeting_minutes.md", title=title, sections=sections, confirmed=True),
        ]
        email = self.create_email_draft(root / "followup.eml", subject=f"Follow-up: {title}", body=self._sections_markdown(title, sections), to=recipients or [], attachments=[item["path"] for item in artifacts[:3]], confirmed=True)
        calendar = self.create_calendar_event(root / "followup.ics", summary=f"Follow-up: {title}", start=meeting.get("followup_start") or (datetime.now(timezone.utc) + timedelta(days=7)).isoformat(), duration_minutes=30, attendees=recipients or [], description="Review meeting actions and blockers.", confirmed=True)
        result = {"schema": "clawos.office-meeting-package.v1", "title": title, "artifacts": artifacts + [email, calendar], "quality_gate": "pass" if all(item.get("valid") for item in artifacts + [email, calendar]) else "revision_required"}
        self._save_run("office_meeting_package", title, result)
        return result

    def create_email_draft(self, output_path: str | Path, *, subject: str, body: str, to: list[str], cc: list[str] | None = None, attachments: list[str] | None = None, confirmed: bool = False) -> dict[str, Any]:
        if confirmed is not True:
            raise PermissionError("confirmed=true is required to write an email draft")
        self._validate_header(subject)
        recipients = [self._email(item) for item in to]
        cc_rows = [self._email(item) for item in (cc or [])]
        message = EmailMessage(policy=SMTP)
        message["Subject"] = subject
        message["To"] = ", ".join(recipients)
        if cc_rows:
            message["Cc"] = ", ".join(cc_rows)
        message["Date"] = datetime.now(timezone.utc)
        message["Message-ID"] = f"<{uuid.uuid4().hex}@clawos.local>"
        message.set_content(body, charset="utf-8")
        for raw_path in attachments or []:
            path = self._path(raw_path)
            if not path.is_file():
                raise ValueError(f"attachment does not exist: {path}")
            mime, _ = mimetypes.guess_type(path.name)
            main, sub = (mime or "application/octet-stream").split("/", 1)
            message.add_attachment(path.read_bytes(), maintype=main, subtype=sub, filename=path.name)
        output = self._path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(message.as_bytes(policy=SMTP))
        result = self.inspect(output)
        result.update({"draft": True, "sent": False, "recipient_count": len(recipients) + len(cc_rows)})
        return result

    def create_calendar_event(self, output_path: str | Path, *, summary: str, start: str, duration_minutes: int = 60, attendees: list[str] | None = None, description: str = "", location: str = "", organizer: str = "", alarm_minutes: int = 15, confirmed: bool = False) -> dict[str, Any]:
        if confirmed is not True:
            raise PermissionError("confirmed=true is required to write a calendar event")
        start_dt = datetime.fromisoformat(str(start).replace("Z", "+00:00"))
        if start_dt.tzinfo is None:
            start_dt = start_dt.replace(tzinfo=timezone.utc)
        end_dt = start_dt + timedelta(minutes=max(1, min(int(duration_minutes or 60), 10080)))
        uid = f"{uuid.uuid4().hex}@clawos.local"
        lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//ClawOS//Office Productivity//EN", "CALSCALE:GREGORIAN", "METHOD:PUBLISH", "BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{self._ics_time(datetime.now(timezone.utc))}", f"DTSTART:{self._ics_time(start_dt)}", f"DTEND:{self._ics_time(end_dt)}", f"SUMMARY:{self._ics(summary)}", f"DESCRIPTION:{self._ics(description)}", f"LOCATION:{self._ics(location)}", "SEQUENCE:0", "STATUS:CONFIRMED"]
        if organizer:
            lines.append(f"ORGANIZER:mailto:{self._email(organizer)}")
        lines.extend(f"ATTENDEE;RSVP=TRUE:mailto:{self._email(item)}" for item in (attendees or []))
        if alarm_minutes > 0:
            lines.extend(["BEGIN:VALARM", "ACTION:DISPLAY", f"TRIGGER:-PT{int(alarm_minutes)}M", f"DESCRIPTION:{self._ics(summary)}", "END:VALARM"])
        lines.extend(["END:VEVENT", "END:VCALENDAR", ""])
        output = self._path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text("\r\n".join(self._fold_ics(line) for line in lines), encoding="utf-8", newline="")
        result = self.inspect(output)
        result.update({"draft": True, "published": False, "uid": uid, "start": start_dt.isoformat(), "end": end_dt.isoformat()})
        return result

    def send_email(self, eml_path: str | Path, smtp_config: dict[str, Any], *, confirmed: bool = False, idempotency_key: str = "") -> dict[str, Any]:
        if confirmed is not True or not idempotency_key:
            raise PermissionError("confirmed=true and idempotency_key are required for email delivery")
        path = self._path(eml_path)
        data = path.read_bytes()
        if self.smtp_sender:
            response = self.smtp_sender(data, smtp_config, idempotency_key)
        else:
            host = str(smtp_config.get("host") or "")
            port = int(smtp_config.get("port") or 465)
            self._public_host(host)
            username = str(smtp_config.get("username") or "")
            password = os.environ.get(str(smtp_config.get("password_env") or ""), "")
            if not host or not username or not password:
                raise ValueError("SMTP host, username and password_env secret are required")
            with smtplib.SMTP_SSL(host, port, context=ssl.create_default_context(), timeout=30) as client:
                client.login(username, password)
                response = client.sendmail(username, smtp_config.get("recipients") or [], data)
        return {"sent": True, "idempotency_key": idempotency_key, "provider_response": response or {}, "sha256": self._sha(path)}

    def publish_calendar(self, ics_path: str | Path, config: dict[str, Any], *, confirmed: bool = False, idempotency_key: str = "") -> dict[str, Any]:
        if confirmed is not True or not idempotency_key:
            raise PermissionError("confirmed=true and idempotency_key are required for calendar publishing")
        path = self._path(ics_path)
        if self.calendar_sender:
            response = self.calendar_sender(path.read_bytes(), config, idempotency_key)
        else:
            url = str(config.get("url") or "")
            parsed = urllib.parse.urlparse(url)
            if parsed.scheme != "https" or not parsed.hostname:
                raise ValueError("CalDAV URL must use public HTTPS")
            self._public_host(parsed.hostname)
            username = str(config.get("username") or "")
            password = os.environ.get(str(config.get("password_env") or ""), "")
            token = __import__("base64").b64encode(f"{username}:{password}".encode()).decode()
            request = urllib.request.Request(url, data=path.read_bytes(), method="PUT", headers={"Content-Type": "text/calendar; charset=utf-8", "Authorization": f"Basic {token}", "If-None-Match": "*", "X-ClawOS-Idempotency-Key": idempotency_key})
            with urllib.request.urlopen(request, timeout=30) as response_obj:
                response = {"status": response_obj.status, "location": response_obj.headers.get("Location", "")}
        return {"published": True, "idempotency_key": idempotency_key, "provider_response": response or {}, "sha256": self._sha(path)}

    def _finalize(self, operation: str, path: Path, engine: str, metadata: dict[str, Any]) -> dict[str, Any]:
        inspection = self.inspect(path)
        visual_quality: dict[str, Any] = {}
        semantic_quality: dict[str, Any] = {}
        if path.suffix.lower().lstrip(".") in {"docx", "xlsx", "pptx", "pdf"}:
            try:
                visual_quality = self.visual_quality.audit(path, strict_visual=False)
                visual_quality["render"] = self._persist_previews(path, visual_quality.get("render") or {})
            except Exception as exc:
                visual_quality = {
                    "schema": "clawos.office-visual-quality.v1",
                    "path": str(path),
                    "quality_gate": "revision_required",
                    "visual_status": "audit_failed",
                    "hard_failures": ["visual_audit_error"],
                    "warnings": [],
                    "error": f"{type(exc).__name__}: {exc}",
                    "render": {"available": False, "pages": 0, "images": []},
                }
        if path.suffix.lower() in {".docx", ".pdf"}:
            semantic_quality = self._semantic_generation_audit(path, metadata.get("semantic_requirements") or {})
        quality_passed = (
            inspection["valid"] and inspection["bytes"] > 0
            and visual_quality.get("quality_gate", "pass") == "pass"
            and semantic_quality.get("gate", "pass") == "pass"
        )
        inspection.update({"operation": operation, "engine": engine, "metadata": metadata,
                           "visual_quality": visual_quality, "semantic_quality": semantic_quality,
                           "quality_gate": "pass" if quality_passed else "revision_required"})
        self._save_run("office_artifact", f"{operation}: {path.name}", inspection)
        return inspection

    @staticmethod
    def _semantic_requirements(sections: list[dict[str, Any]], metadata: dict[str, Any], fmt: str) -> dict[str, int]:
        if fmt not in {"docx", "pdf"}:
            return {}
        return {
            "headings": sum(bool(item.get("heading")) for item in sections),
            "tables": sum(bool(item.get("table")) for item in sections),
            "figures": sum(len(item.get("figures") or []) for item in sections),
            "formulas": sum(len(item.get("formulas") or []) for item in sections),
            "footnotes": len(metadata.get("footnotes") or []),
            "comments": len(metadata.get("comments") or []) if fmt == "docx" else 0,
            "revisions": len(metadata.get("revisions") or []) if fmt == "docx" else 0,
        }

    @staticmethod
    def _semantic_generation_audit(path: Path, requirements: dict[str, int]) -> dict[str, Any]:
        try:
            from .document_semantic_graph_service import DocumentSemanticGraphService

            graph = DocumentSemanticGraphService().parse(path)
            actual = {
                "headings": len(graph.get("sections") or []), "tables": len(graph.get("tables") or []),
                "figures": len(graph.get("figures") or []), "formulas": len(graph.get("formulas") or []),
                "footnotes": len(graph.get("footnotes") or []), "comments": len(graph.get("comments") or []),
                "revisions": len(graph.get("revisions") or []),
            }
            checks = {key: actual.get(key, 0) >= int(value or 0) for key, value in requirements.items()}
            return {
                "schema": "clawos.office-semantic-generation-audit.v1", "requirements": requirements,
                "actual": actual, "checks": checks, "gate": "pass" if all(checks.values()) else "revision_required",
                "parser_quality": graph.get("quality") or {},
            }
        except Exception as exc:
            return {"schema": "clawos.office-semantic-generation-audit.v1", "requirements": requirements,
                    "actual": {}, "checks": {}, "gate": "revision_required",
                    "error": f"{type(exc).__name__}: {exc}"}

    def _persist_previews(self, source: Path, render: dict[str, Any]) -> dict[str, Any]:
        if not self.preview_root or not render.get("available"):
            return render
        images = [Path(item) for item in render.get("images") or [] if Path(item).is_file()]
        if not images:
            return render
        digest = self._sha(source)[:16]
        target_root = self.preview_root / digest
        target_root.mkdir(parents=True, exist_ok=True)
        persisted = []
        for index, image in enumerate(images[:20], 1):
            target = target_root / f"page-{index:03d}{image.suffix.lower()}"
            shutil.copyfile(image, target)
            persisted.append(str(target))
        return {**render, "images": persisted, "persisted": True}

    def _save_run(self, run_type: str, title: str, result: dict[str, Any]) -> None:
        if not self.store:
            return
        run = RunRecord(run_id=f"{run_type}-{int(time.time())}-{uuid.uuid4().hex[:6]}", run_type=run_type, title=title, status="passed" if result.get("quality_gate", "pass") == "pass" else "failed", actor="office-agent", scenario="office", completed_at=now_iso(), outputs=result, metrics={"artifact_count": len(result.get("artifacts") or []) or 1}, tags=["office", "artifact", "quality-gated"])
        for item in result.get("artifacts") or [result]:
            if item.get("path"):
                run.add_artifact(item.get("format", "office"), Path(item["path"]).name, item["path"], sha256=item.get("sha256"), engine=item.get("engine"))
        self.store.save_run_record(run.to_dict())

    def _path(self, value: str | Path) -> Path:
        path = Path(value).expanduser().resolve()
        if not any(self._within(path, root) for root in self.allowed_roots):
            raise ValueError(f"path is outside allowed roots: {path}")
        return path

    @staticmethod
    def _within(path: Path, root: Path) -> bool:
        try:
            path.relative_to(root)
            return True
        except ValueError:
            return False

    @staticmethod
    def _sha(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _module_engine(module: str, professional: str, fallback: str) -> str:
        try:
            __import__(module)
            return professional
        except ImportError:
            return fallback

    @staticmethod
    def _openxml_fallback(path: Path, required_part: str, original_error: Exception) -> dict[str, Any]:
        try:
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
                if "[Content_Types].xml" not in names or required_part not in names:
                    raise ValueError("required OpenXML parts are missing")
                bad = archive.testzip()
                if bad:
                    raise ValueError(f"corrupt OpenXML part: {bad}")
                return {"openxml_parts": len(names), "required_part": required_part, "fallback_validation": True, "professional_engine_unavailable": type(original_error).__name__}
        except Exception as exc:
            return {"open_error": f"{type(exc).__name__}: {exc}", "professional_engine_error": f"{type(original_error).__name__}: {original_error}"}

    @staticmethod
    def _sections_markdown(title: str, sections: list[dict[str, Any]]) -> str:
        lines = [f"# {title}", ""]
        for item in sections:
            if item.get("heading"):
                lines.extend([f"## {item['heading']}", ""])
            lines.extend([str(value) for value in item.get("paragraphs") or []])
            lines.extend([f"- {value}" for value in item.get("bullets") or []])
            if item.get("table"):
                table = item["table"]
                rows = table.get("rows", []) if isinstance(table, dict) else table
                if rows:
                    width = max(len(row) for row in rows)
                    def render_row(row):
                        cells = [str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>") for value in row]
                        return "| " + " | ".join(cells + [""] * (width-len(cells))) + " |"
                    lines.extend(["", render_row(rows[0]), "| " + " | ".join(["---"]*width) + " |"])
                    lines.extend(render_row(row) for row in rows[1:])
            lines.append("")
        return "\n".join(lines).strip() + "\n"

    @staticmethod
    def _write_csv(path: Path, rows: list[Any]) -> None:
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            for row in rows:
                values = row.values() if isinstance(row, dict) else row
                writer.writerow([OfficeProductivityService._safe_csv(value) for value in values])

    @staticmethod
    def _safe_csv(value: Any) -> Any:
        text = str(value) if value is not None else ""
        return "'" + text if text[:1] in {"=", "+", "-", "@"} else value

    @staticmethod
    def _validate_header(value: str) -> None:
        if "\r" in value or "\n" in value:
            raise ValueError("email header injection blocked")

    @staticmethod
    def _email(value: str) -> str:
        OfficeProductivityService._validate_header(str(value))
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", str(value or "")):
            raise ValueError(f"invalid email address: {value}")
        return str(value)

    @staticmethod
    def _ics(value: str) -> str:
        return str(value or "").replace("\\", "\\\\").replace("\r\n", "\\n").replace("\n", "\\n").replace(",", "\\,").replace(";", "\\;")

    @staticmethod
    def _ics_time(value: datetime) -> str:
        return value.astimezone(timezone.utc).strftime("%Y%m%dT%H%M%SZ")

    @staticmethod
    def _fold_ics(line: str) -> str:
        chunks = []
        current = ""
        for char in line:
            if len((current + char).encode("utf-8")) > 73:
                chunks.append(current)
                current = " " + char
            else:
                current += char
        chunks.append(current)
        return "\r\n".join(chunks)

    @staticmethod
    def _public_host(host: str) -> None:
        if not host:
            raise ValueError("host is required")
        for item in socket.getaddrinfo(host, None):
            ip = ipaddress.ip_address(item[4][0])
            if not ip.is_global:
                raise ValueError("private or unsafe connector host is not allowed")
