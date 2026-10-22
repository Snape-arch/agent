"""Bounded, reproducible tabular processing without arbitrary code execution."""
import csv
import hashlib
import io
import json
import math
import re
import statistics
import uuid
import zipfile
from collections import Counter, defaultdict
from fastapi import APIRouter, HTTPException, Request, UploadFile, File
from fastapi.responses import Response

MAX_ROWS, MAX_COLUMNS, MAX_BYTES = 20000, 100, 15*1024*1024


def number(value):
    if type(value) in (int, float):
        try:return float(value) if math.isfinite(value) and abs(value)<=1e100 else None
        except OverflowError:return None
    if isinstance(value, str) and re.fullmatch(r'[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?', value.strip()):
        try:
            n=float(value)
            return n if math.isfinite(n) and abs(n)<=1e100 else None
        except (ValueError,OverflowError):pass
    return None


def missing(v):
    return v is None or (isinstance(v,str) and not v.strip())


def validate(columns, rows):
    if not isinstance(columns,list) or not 1<=len(columns)<=MAX_COLUMNS or any(not isinstance(x,str) or not x.strip() or len(x)>120 for x in columns) or len(set(columns))!=len(columns):
        raise ValueError('列名需要非空且唯一，最多 100 列')
    if not isinstance(rows,list) or len(rows)>MAX_ROWS:
        raise ValueError('最多处理 20,000 行')
    for r in rows:
        if not isinstance(r,dict) or set(r)!=set(columns):raise ValueError('每行需包含完整列')
        for v in r.values():
            if v is not None and type(v) not in (str,int,float,bool):raise ValueError('单元格仅支持文字、数值、布尔或空值')
            if isinstance(v,str) and len(v)>20000:raise ValueError('单元格内容过长')
            if type(v) is float and not math.isfinite(v):raise ValueError('不支持非有限数值')
            if type(v) in (int,float) and abs(v)>1e100:raise ValueError('数值绝对值需不超过 1e100')
    if len(json.dumps(rows,ensure_ascii=False).encode())>MAX_BYTES:raise ValueError('数据展开后超过 15 MB')


def parse_upload(filename, raw):
    if not raw or len(raw)>MAX_BYTES:raise ValueError('上传需要 1 字节至 15 MB')
    warnings=[]
    if filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist())>40*1024*1024:raise ValueError('工作簿展开后过大')
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
            try:
                if len(wb.sheetnames)>1:warnings.append('已导入首个工作表；其他工作表未参与处理。')
                ws=wb.worksheets[0]
                if ws.max_row>MAX_ROWS+1 or ws.max_column>MAX_COLUMNS:raise ValueError('工作表超过 20,000 行 / 100 列')
                matrix=[];formulas=0
                for row in ws.iter_rows():
                    values=[]
                    for cell in row:
                        if cell.data_type=='f':formulas+=1;values.append(None)
                        else:
                            v=cell.value
                            if v is not None and type(v) not in (str,int,float,bool):v=v.isoformat() if hasattr(v,'isoformat') else str(v)
                            values.append(v)
                    matrix.append(values)
                    if len(matrix)>MAX_ROWS+1:raise ValueError('工作表行数过大')
                if formulas:warnings.append(f'{formulas} 个公式单元格作为空值导入；本工具不执行公式，请先在表格软件计算并转为值。')
            finally:wb.close()
        except (zipfile.BadZipFile,KeyError):raise ValueError('不是有效的 XLSX 文件')
    elif filename.lower().endswith(('.csv','.tsv')):
        for encoding in ('utf-8-sig','gb18030'):
            try: text=raw.decode(encoding);break
            except UnicodeDecodeError:continue
        else:raise ValueError('仅支持 UTF-8 / GB18030 编码')
        if '\x00' in text:raise ValueError('输入含二进制内容')
        try: dialect=csv.Sniffer().sniff(text[:8000],delimiters=',\t;')
        except csv.Error: dialect=csv.excel_tab if filename.lower().endswith('.tsv') else csv.excel
        matrix=[]
        try:
            for row in csv.reader(io.StringIO(text),dialect):
                if row:matrix.append(row)
                if len(matrix)>MAX_ROWS+1:raise ValueError('最多 20,000 行')
        except csv.Error:raise ValueError('CSV 格式无效')
    else:raise ValueError('请选择 CSV、TSV 或 XLSX')
    if not matrix:raise ValueError('文件没有表头')
    columns=[str(x or '').strip() for x in matrix[0]]
    if any(len(r)!=len(columns) for r in matrix[1:]):raise ValueError('数据行与表头列数不一致')
    rows=[dict(zip(columns,r)) for r in matrix[1:]]
    validate(columns,rows)
    return columns,rows,warnings


def profile(columns,rows):
    summaries=[]
    for key in columns:
        values=[r[key] for r in rows if not missing(r[key])]
        nums=[number(v) for v in values];valid=[n for n in nums if n is not None]
        numeric=bool(values) and len(valid)==len(values)
        summary=dict(column=key,type='number' if numeric else 'text',missing=len(rows)-len(values),unique=len(set(map(str,values))))
        if numeric:
            if len(valid)>1:q1,_,q3=statistics.quantiles(valid,n=4,method='inclusive')
            else:q1=q3=valid[0]
            iqr=q3-q1
            summary.update(min=min(valid),max=max(valid),mean=statistics.fmean(valid),median=statistics.median(valid),std=statistics.pstdev(valid),outliers=sum(n<q1-1.5*iqr or n>q3+1.5*iqr for n in valid))
        summaries.append(summary)
    duplicates=len(rows)-len({json.dumps(r,sort_keys=True,ensure_ascii=False) for r in rows})
    return dict(rows=len(rows),columns=len(columns),missing=sum(x['missing'] for x in summaries),duplicates=duplicates,fields=summaries)


def process(columns,rows,recipe):
    if not isinstance(recipe,list) or len(recipe)>20:raise ValueError('处理流程最多 20 步')
    columns=list(columns);rows=[dict(r) for r in rows];steps=[]
    for spec in recipe:
        if not isinstance(spec,dict):raise ValueError('处理步骤无效')
        op=spec.get('op');before=len(rows);key=spec.get('column')
        if op in ('trim','fill','numeric','filter') and key not in columns:raise ValueError('处理列不存在')
        if op=='trim':
            for r in rows:
                if isinstance(r[key],str):r[key]=r[key].strip()
        elif op=='numeric':
            invalid=sum(not missing(r[key]) and number(r[key]) is None for r in rows)
            if invalid:raise ValueError(f'{key} 有 {invalid} 个非数值，请先过滤或修正；不会静默丢弃。')
            for r in rows:r[key]=number(r[key])
        elif op=='fill':
            strategy=spec.get('strategy','constant');value=spec.get('value')
            if strategy in ('mean','median'):
                nums=[number(r[key]) for r in rows if not missing(r[key])]
                if not nums or any(n is None for n in nums):raise ValueError('均值/中位数填补需要有效数值列')
                value=statistics.fmean(nums) if strategy=='mean' else statistics.median(nums)
            elif strategy!='constant':raise ValueError('填补方式无效')
            if value is None:raise ValueError('请指定填补值')
            for r in rows:
                if missing(r[key]):r[key]=value
        elif op=='deduplicate':
            keys=spec.get('columns') or columns
            if not isinstance(keys,list) or any(k not in columns for k in keys):raise ValueError('去重列不存在')
            seen=set();kept=[]
            for r in rows:
                item=json.dumps([r[k] for k in keys],ensure_ascii=False)
                if item not in seen:seen.add(item);kept.append(r)
            rows=kept
        elif op=='filter':
            cmp=spec.get('comparison');value=spec.get('value')
            if cmp not in ('eq','ne','gt','ge','lt','le','contains','present'):raise ValueError('筛选条件无效')
            if cmp in ('gt','ge','lt','le') and number(value) is None:raise ValueError('数值筛选需要数值阈值')
            def keep(r):
                v=r[key]
                if cmp=='present':return not missing(v)
                if cmp=='contains':return str(value) in str(v or '')
                if cmp=='eq':return str(v)==str(value)
                if cmp=='ne':return str(v)!=str(value)
                n=number(v);threshold=number(value)
                if n is None:return False
                return {'gt':n>threshold,'ge':n>=threshold,'lt':n<threshold,'le':n<=threshold}[cmp]
            rows=[r for r in rows if keep(r)]
        elif op=='derive':
            a,b,name,method=spec.get('left'),spec.get('right'),spec.get('name'),spec.get('method')
            if a not in columns or b not in columns or not isinstance(name,str) or not 1<=len(name)<=120 or name in columns or method not in ('add','subtract','multiply','divide'):raise ValueError('派生列参数无效')
            columns.append(name)
            for r in rows:
                x,y=number(r[a]),number(r[b]);n=None
                if x is not None and y is not None:
                    if method=='divide':n=x/y if y else None
                    else:n={'add':lambda:x+y,'subtract':lambda:x-y,'multiply':lambda:x*y}[method]()
                r[name]=n if n is not None and math.isfinite(n) else None
        elif op=='group':
            group=spec.get('group');metric=spec.get('metric');agg=spec.get('aggregate','sum')
            if group not in columns or metric not in columns or group==metric or agg not in ('sum','mean','min','max','count'):raise ValueError('分组参数无效')
            buckets=defaultdict(list)
            for r in rows:buckets[str(r[group]) if not missing(r[group]) else '(空)'].append(r[metric])
            if agg!='count' and any(number(v) is None for values in buckets.values() for v in values if not missing(v)):raise ValueError('聚合指标含非数值')
            out='value';columns=[group,out,'count'] if group not in ('value','count') else ['group',out,'count']
            grouped=[]
            for label,values in sorted(buckets.items()):
                ns=[number(v) for v in values if number(v) is not None]
                val=len(values) if agg=='count' else (sum(ns) if agg=='sum' else statistics.fmean(ns) if agg=='mean' and ns else min(ns) if agg=='min' and ns else max(ns) if agg=='max' and ns else None)
                grouped.append(dict(zip(columns,[label,val,len(values)])))
            rows=grouped
        elif op=='sort':
            if key not in columns:raise ValueError('排序列不存在')
            numeric=all(number(r[key]) is not None for r in rows if not missing(r[key]))
            blanks=[r for r in rows if missing(r[key])]
            rows=sorted([r for r in rows if not missing(r[key])],key=lambda r:number(r[key]) if numeric else str(r[key]),reverse=spec.get('descending') is True)+blanks
        else:raise ValueError('不支持的处理步骤')
        validate(columns,rows)
        steps.append(dict(step=spec,before=before,after=len(rows)))
    return columns,rows,steps


def chart(columns,rows,spec):
    kind,x,y=spec.get('kind','bar'),spec.get('x'),spec.get('y')
    if x not in columns or (kind!='histogram' and y not in columns):raise ValueError('图表列不存在')
    if kind not in ('bar','line','scatter','histogram'):raise ValueError('图表类型无效')
    if kind=='histogram':
        values=[number(r[x]) for r in rows if number(r[x]) is not None]
        if not values:raise ValueError('分布图需要数值列')
        bins=12;lo,hi=min(values),max(values);width=(hi-lo)/bins if hi!=lo else 1;counts=[0]*bins
        for n in values:counts[min(bins-1,int((n-lo)/width))]+=1
        points=[dict(x=f'{lo+i*width:.3g}–{lo+(i+1)*width:.3g}',y=n) for i,n in enumerate(counts)]
        return dict(kind=kind,x=x,y='频数',points=points,omitted=len(rows)-len(values),method='12 个等宽区间')
    if kind=='scatter':
        points=[dict(x=number(r[x]),y=number(r[y])) for r in rows if number(r[x]) is not None and number(r[y]) is not None]
        n=len(points);corr=None
        if n>1:
            mx=statistics.fmean(p['x'] for p in points);my=statistics.fmean(p['y'] for p in points)
            xx=sum((p['x']-mx)**2 for p in points);yy=sum((p['y']-my)**2 for p in points)
            if xx and yy:corr=sum((p['x']-mx)*(p['y']-my) for p in points)/math.sqrt(xx*yy)
        return dict(kind=kind,x=x,y=y,points=points[:2000],omitted=len(rows)-n,display_omitted=max(0,n-2000),correlation=corr,method='Pearson 相关仅表示线性关联，不表示因果；最多绘制前 2000 个有效点')
    groups=defaultdict(list)
    for r in rows:
        n=number(r[y])
        if n is not None:groups[str(r[x]) if not missing(r[x]) else '(空)'].append(n)
    agg=spec.get('aggregate','sum')
    if agg not in ('sum','mean','count'):raise ValueError('图表聚合方式无效')
    if not groups:raise ValueError('图表需要有效数值')
    keys=sorted(groups,key=lambda k:number(k)) if kind=='line' and all(number(k) is not None for k in groups) else sorted(groups)
    points=[dict(x=k,y=sum(groups[k]) if agg=='sum' else statistics.fmean(groups[k]) if agg=='mean' else len(groups[k])) for k in keys]
    return dict(kind=kind,x=x,y=y,points=points[:100],omitted=sum(number(r[y]) is None for r in rows),display_omitted=max(0,len(points)-100),method=f'按 {x} 分组，{agg}；最多显示前 100 组')


def safe_cell(value):
    # Prevent formulas/DDE when exported text is opened in Excel.
    return "'"+value if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r')) else value


class DataLab:
    def __init__(self,store):
        self.store=store
        with store._conn() as c:
            c.executescript('''CREATE TABLE IF NOT EXISTS data_lab_datasets(id TEXT PRIMARY KEY,name TEXT NOT NULL,version INTEGER NOT NULL,source TEXT NOT NULL,result TEXT NOT NULL,recipe TEXT NOT NULL,source_hash TEXT NOT NULL,warnings TEXT NOT NULL,updated_at TEXT DEFAULT (datetime('now')));
            CREATE TABLE IF NOT EXISTS data_lab_runs(id TEXT PRIMARY KEY,dataset_id TEXT NOT NULL,version INTEGER NOT NULL,recipe TEXT NOT NULL,metrics TEXT NOT NULL,created_at TEXT DEFAULT (datetime('now')));''')

    def create(self,name,columns,rows,warnings=None):
        validate(columns,rows)
        if not isinstance(name,str) or not 1<=len(name.strip())<=200:raise ValueError('名称需要 1–200 字')
        nid=uuid.uuid4().hex;payload=json.dumps(dict(columns=columns,rows=rows),ensure_ascii=False,allow_nan=False)
        with self.store._conn() as c:c.execute('INSERT INTO data_lab_datasets(id,name,version,source,result,recipe,source_hash,warnings) VALUES (?,?,1,?,?,?, ?,?)',(nid,name.strip(),payload,payload,'[]',hashlib.sha256(payload.encode()).hexdigest(),json.dumps(warnings or [],ensure_ascii=False)))
        return self.get(nid)

    def get(self,nid,full=False):
        with self.store._conn() as c:row=c.execute('SELECT * FROM data_lab_datasets WHERE id=?',(nid,)).fetchone()
        if not row:raise HTTPException(404,'数据集不存在')
        r=dict(row);result=json.loads(r.pop('result'));r.pop('source');r['recipe']=json.loads(r['recipe']);r['warnings']=json.loads(r['warnings']);r['profile']=profile(**result);r['columns']=result['columns'];r['rows']=result['rows'] if full else result['rows'][:100];r['preview_limit']=100
        return r

    def run(self,nid,version,recipe):
        with self.store._conn() as c:
            c.execute('BEGIN IMMEDIATE');r=c.execute('SELECT * FROM data_lab_datasets WHERE id=?',(nid,)).fetchone()
            if not r:raise HTTPException(404,'数据集不存在')
            if type(version) is not int or version!=r['version']:raise HTTPException(409,'数据版本已变化，请刷新后处理')
            source=json.loads(r['source']);columns,rows,steps=process(source['columns'],source['rows'],recipe)
            result=json.dumps(dict(columns=columns,rows=rows),ensure_ascii=False,allow_nan=False);recipe_text=json.dumps(recipe,ensure_ascii=False)
            c.execute("UPDATE data_lab_datasets SET result=?,recipe=?,version=version+1,updated_at=datetime('now') WHERE id=?",(result,recipe_text,nid))
            runid=uuid.uuid4().hex
            c.execute('INSERT INTO data_lab_runs(id,dataset_id,version,recipe,metrics) VALUES (?,?,?,?,?)',(runid,nid,version+1,recipe_text,json.dumps(dict(source_hash=r['source_hash'],source_rows=len(source['rows']),result_rows=len(rows),steps=steps),ensure_ascii=False)))
        return dict(**self.get(nid),run_id=runid,steps=steps)


def data_lab_router(factory,access_factory):
    router=APIRouter(prefix='/api/data-lab')
    def guard(request):
        if access_factory().identity(request)['id']!='local-user':raise HTTPException(403,'数据工作台仅向本机所有者开放')
    @router.get('')
    def index(request:Request):
        guard(request)
        with factory().store._conn() as c:return dict(datasets=[dict(r) for r in c.execute('SELECT id,name,version,updated_at FROM data_lab_datasets ORDER BY updated_at DESC')])
    @router.post('/upload')
    async def upload(request:Request,file:UploadFile=File(...)):
        guard(request)
        try:
            from starlette.concurrency import run_in_threadpool
            raw=await file.read(MAX_BYTES+1)
            cols,rows,warnings=await run_in_threadpool(parse_upload,file.filename or '',raw)
            return await run_in_threadpool(factory().create,(file.filename or '数据集')[:200],cols,rows,warnings)
        except (ValueError,TypeError) as e:raise HTTPException(400,str(e))
    @router.post('')
    async def create(request:Request):
        guard(request)
        try:
            if len(await request.body())>MAX_BYTES:raise ValueError('数据过大')
            from starlette.concurrency import run_in_threadpool
            d=await request.json();return await run_in_threadpool(factory().create,d.get('name'),d.get('columns'),d.get('rows'))
        except (ValueError,TypeError) as e:raise HTTPException(400,str(e))
    @router.get('/{nid}')
    def get(nid:str,request:Request):guard(request);return factory().get(nid)
    @router.post('/{nid}/process')
    async def run(nid:str,request:Request):
        guard(request)
        try:
            from starlette.concurrency import run_in_threadpool
            d=await request.json();return await run_in_threadpool(factory().run,nid,d.get('version'),d.get('recipe'))
        except (ValueError,TypeError,OverflowError) as e:raise HTTPException(400,str(e))
    @router.post('/{nid}/chart')
    async def draw(nid:str,request:Request):
        guard(request)
        try:
            from starlette.concurrency import run_in_threadpool
            d=await run_in_threadpool(factory().get,nid,True)
            return await run_in_threadpool(chart,d['columns'],d['rows'],await request.json())
        except (ValueError,TypeError,OverflowError) as e:raise HTTPException(400,str(e))
    @router.get('/{nid}/runs')
    def runs(nid:str,request:Request):
        guard(request);s=factory();s.get(nid)
        with s.store._conn() as c:return dict(runs=[dict(r) for r in c.execute('SELECT * FROM data_lab_runs WHERE dataset_id=? ORDER BY version DESC LIMIT 50',(nid,))])
    @router.get('/{nid}/export')
    def export(nid:str,request:Request,format:str='csv'):
        guard(request);d=factory().get(nid,True)
        if format=='recipe':return Response(json.dumps(dict(source_hash=d['source_hash'],version=d['version'],recipe=d['recipe']),ensure_ascii=False,indent=2),media_type='application/json',headers={'Content-Disposition':'attachment; filename="recipe.json"'})
        if format=='csv':
            buf=io.StringIO();writer=csv.writer(buf);writer.writerow([safe_cell(c) for c in d['columns']]);writer.writerows([[safe_cell(r[c]) for c in d['columns']] for r in d['rows']]);raw=buf.getvalue().encode('utf-8-sig');mime='text/csv'
        elif format=='xlsx':
            from openpyxl import Workbook
            wb=Workbook();ws=wb.active;ws.title='Data';ws.append([safe_cell(c) for c in d['columns']])
            for r in d['rows']:ws.append([safe_cell(r[c]) for c in d['columns']])
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            meta=wb.create_sheet('Recipe');meta.append(['Source SHA-256',d['source_hash']]);meta.append(['Version',d['version']]);meta.append(['Recipe',json.dumps(d['recipe'],ensure_ascii=False)])
            buf=io.BytesIO();wb.save(buf);raw=buf.getvalue();mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        else:raise HTTPException(400,'导出格式无效')
        return Response(raw,media_type=mime,headers={'Content-Disposition':f'attachment; filename="dataset-{nid}.{format}"'})
    return router
