from __future__ import annotations

import json
import math
import shutil
import statistics
import subprocess
import tempfile
import os
import re
import zipfile
from xml.etree import ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any


class OfficeVisualQualityService:
    """Render-aware office QA plus auditable tabular analysis/workbook generation."""

    def audit(self, file_path: str | Path, *, template_path: str | Path = "", strict_visual: bool = False) -> dict[str, Any]:
        path = Path(file_path).expanduser().resolve()
        if not path.is_file():
            raise ValueError("file_path must exist")
        fmt = path.suffix.lower().lstrip(".")
        checks: list[dict[str, Any]] = []
        render = {"available": False, "engine": "", "pages": 0, "images": []}
        if fmt == "pdf":
            checks.extend(self._pdf_checks(path))
            render = self._render_pdf(path)
        elif fmt == "docx":
            checks.extend(self._docx_checks(path))
            render = self._render_via_office(path)
            checks.extend(self.rendered_note_checks(path, render.get("pdf", "")))
        elif fmt == "pptx":
            checks.extend(self._pptx_checks(path))
            render = self._render_via_office(path)
        elif fmt == "xlsx":
            checks.extend(self._xlsx_checks(path))
            render = self._render_via_office(path)
        else:
            raise ValueError(f"visual QA is not supported for {fmt}")
        if template_path:
            checks.append(self._template_fidelity(path, Path(template_path).expanduser().resolve()))
        if render.get("stderr") and "No display font" in render["stderr"]:
            # TeXLive's Windows Poppler build emits this for its optional
            # ArialUnicode fallback even when the PDF embeds its own font and
            # every page renders. Preserve the diagnostic, but only block when
            # rendering itself failed.
            checks.append({"name": "render_font_resolution", "passed": bool(render.get("available") and render.get("pages")), "severity": "high", "detail": render["stderr"]})
        visual_gate = render["available"] and render["pages"] > 0
        checks.append({"name": "rendered_pages", "passed": visual_gate, "severity": "high" if strict_visual else "warning", "detail": render})
        hard_failures = [item for item in checks if not item["passed"] and item.get("severity") == "high"]
        warnings = [item for item in checks if not item["passed"] and item.get("severity") != "high"]
        return {
            "schema": "clawos.office-visual-quality.v1", "path": str(path), "format": fmt,
            "quality_gate": "pass" if not hard_failures else "revision_required",
            "visual_status": "rendered" if visual_gate else "structural_only",
            "strict_visual": strict_visual, "checks": checks,
            "hard_failures": [item["name"] for item in hard_failures],
            "warnings": [item["name"] for item in warnings], "render": render,
        }

    def analyze_and_build(self, rows: list[dict[str, Any]], output_path: str | Path, *, title: str = "Data Analysis", group_by: str = "", metrics: list[str] | None = None) -> dict[str, Any]:
        if not rows:
            raise ValueError("rows must be a non-empty list of objects")
        columns = list(dict.fromkeys(key for row in rows for key in row))
        numeric = [key for key in columns if any(isinstance(row.get(key), (int, float)) and not isinstance(row.get(key), bool) for row in rows)]
        selected = [item for item in (metrics or numeric) if item in numeric]
        summaries = {}
        for key in selected:
            values = [float(row[key]) for row in rows if isinstance(row.get(key), (int, float)) and not isinstance(row.get(key), bool)]
            summaries[key] = {"count": len(values), "sum": sum(values), "mean": statistics.fmean(values), "median": statistics.median(values), "min": min(values), "max": max(values), "stdev": statistics.stdev(values) if len(values) > 1 else 0.0}
        grouped: dict[str, dict[str, Any]] = {}
        if group_by and group_by in columns:
            buckets: dict[str, list[dict[str, Any]]] = defaultdict(list)
            for row in rows:
                buckets[str(row.get(group_by, ""))].append(row)
            for name, items in buckets.items():
                grouped[name] = {key: sum(float(row.get(key) or 0) for row in items if isinstance(row.get(key), (int, float))) for key in selected}
                grouped[name]["count"] = len(items)
        correlations = {}
        for i, left in enumerate(selected):
            for right in selected[i + 1:]:
                pairs = [(float(row[left]), float(row[right])) for row in rows if isinstance(row.get(left), (int, float)) and isinstance(row.get(right), (int, float))]
                if len(pairs) >= 2:
                    xs, ys = zip(*pairs)
                    dx, dy = statistics.pstdev(xs), statistics.pstdev(ys)
                    correlations[f"{left}|{right}"] = 0.0 if not dx or not dy else sum((x-statistics.fmean(xs))*(y-statistics.fmean(ys)) for x,y in pairs)/(len(pairs)*dx*dy)
        output = Path(output_path).expanduser().resolve()
        self._build_workbook(rows, output, title, columns, selected, summaries, grouped, group_by)
        audit = self.audit(output)
        return {"schema": "clawos.data-analysis.v1", "rows": len(rows), "columns": columns, "numeric_metrics": selected, "summary": summaries, "grouped": grouped, "correlations": correlations, "workbook": str(output), "quality": audit}

    @staticmethod
    def _docx_checks(path: Path) -> list[dict[str, Any]]:
        try:
            from docx import Document
            from docx.oxml.ns import qn
            doc = Document(path)
            texts = [p.text for p in doc.paragraphs]
            heading_count = sum(bool(p.style and p.style.name.startswith("Heading")) for p in doc.paragraphs)
            table_count = len(doc.tables)
            repeated_headers = all(not table.rows or table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None for table in doc.tables)
        except ImportError:
            import re, zipfile
            with zipfile.ZipFile(path) as archive: xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
            texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml)
            heading_count = len(re.findall(r'w:pStyle w:val="Heading', xml)); table_count = xml.count("<w:tbl>")
            repeated_headers = table_count == 0 or xml.count("<w:tblHeader") >= table_count
        mixed = any(any('\u4e00' <= c <= '\u9fff' for c in text) and any(c.isascii() and c.isalpha() for c in text) for text in texts)
        return [
            {"name": "semantic_styles", "passed": heading_count > 0, "severity": "high", "detail": {"headings": heading_count}},
            {"name": "table_repeat_headers", "passed": repeated_headers, "severity": "high", "detail": {"tables": table_count}},
            {"name": "mixed_language", "passed": mixed, "severity": "warning", "detail": "Chinese/English mixed text present" if mixed else "fixture has no mixed-language paragraph"},
        ]

    @staticmethod
    def _pptx_checks(path: Path) -> list[dict[str, Any]]:
        from pptx import Presentation
        prs = Presentation(path)
        sw, sh = prs.slide_width, prs.slide_height
        overflow, overlaps, small_fonts = [], [], []
        for si, slide in enumerate(prs.slides, 1):
            shapes = [s for s in slide.shapes if getattr(s, "width", 0) and getattr(s, "height", 0)]
            for shape in shapes:
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > sw or shape.top + shape.height > sh:
                    overflow.append({"slide": si, "shape": shape.name})
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.size and run.font.size.pt < 16:
                                small_fonts.append({"slide": si, "shape": shape.name, "pt": run.font.size.pt})
            for i, a in enumerate(shapes):
                for b in shapes[i+1:]:
                    if not (a.left+a.width <= b.left or b.left+b.width <= a.left or a.top+a.height <= b.top or b.top+b.height <= a.top):
                        if getattr(a, "has_text_frame", False) and getattr(b, "has_text_frame", False) and a.text.strip() and b.text.strip():
                            overlaps.append({"slide": si, "a": a.name, "b": b.name})
        return [
            {"name": "slide_canvas_overflow", "passed": not overflow, "severity": "high", "detail": overflow},
            {"name": "text_shape_overlap", "passed": not overlaps, "severity": "high", "detail": overlaps},
            {"name": "minimum_body_font_16pt", "passed": not small_fonts, "severity": "high", "detail": small_fonts},
        ]

    @staticmethod
    def _xlsx_checks(path: Path) -> list[dict[str, Any]]:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False)
        formulas, errors, filters, frozen, styled, charts = 0, [], 0, 0, 0, 0
        for ws in wb.worksheets:
            filters += bool(ws.auto_filter.ref); frozen += bool(ws.freeze_panes); charts += len(ws._charts)
            for row in ws.iter_rows():
                for cell in row:
                    formulas += isinstance(cell.value, str) and cell.value.startswith("=")
                    styled += bool(cell.value is not None and cell.style_id)
                    if isinstance(cell.value, str) and any(token in cell.value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                        errors.append(f"{ws.title}!{cell.coordinate}")
        sheet_count = len(wb.worksheets); wb.close()
        return [
            {"name": "formula_integrity", "passed": not errors, "severity": "high", "detail": {"formulas": formulas, "errors": errors}},
            {"name": "filters", "passed": filters == sheet_count, "severity": "high", "detail": {"sheets": sheet_count, "with_filters": filters}},
            {"name": "freeze_panes", "passed": frozen == sheet_count, "severity": "high", "detail": {"with_freeze": frozen}},
            {"name": "semantic_formatting", "passed": styled > 0, "severity": "high", "detail": {"styled_cells": styled}},
            {"name": "analysis_visual", "passed": charts > 0 or sheet_count > 1, "severity": "warning", "detail": {"charts": charts, "sheets": sheet_count}},
        ]

    @staticmethod
    def _pdf_checks(path: Path) -> list[dict[str, Any]]:
        try:
            from pypdf import PdfReader
            reader = PdfReader(path); chars = [len(page.extract_text() or "") for page in reader.pages]; pages = len(reader.pages)
        except ImportError:
            executable = shutil.which("pdftotext")
            if executable:
                with tempfile.TemporaryDirectory(prefix="clawos-pdf-text-") as td:
                    target = Path(td) / "text.txt"; result = subprocess.run([executable, str(path), str(target)], capture_output=True, timeout=60)
                    text = target.read_text(encoding="utf-8", errors="replace") if target.is_file() else ""
                chunks = text.split("\f"); chars = [len(item.strip()) for item in chunks if item.strip()]; pages = max(1, len(chars))
            else:
                raw=path.read_bytes(); pages=max(1,raw.count(b"/Type /Page")-raw.count(b"/Type /Pages")); chars=[0]*pages
        return [
            {"name": "pdf_pages", "passed": pages > 0, "severity": "high", "detail": {"pages": pages}},
            {"name": "searchable_or_ocr_required", "passed": all(value >= 20 for value in chars), "severity": "high", "detail": {"text_chars_per_page": chars, "ocr_required_pages": [i+1 for i,v in enumerate(chars) if v < 20]}},
        ]

    @staticmethod
    def _render_pdf(path: Path) -> dict[str, Any]:
        # Prefer the native executable on Windows. Some managed runtimes add a
        # `pdftoppm.CMD` shim earlier in PATHEXT; that shim cannot resolve its
        # relative payload when launched from a project virtual environment.
        explicit = Path(os.environ.get("CLAWOS_POPPLER_BIN", "")) / "pdftoppm.exe" if os.environ.get("CLAWOS_POPPLER_BIN") else None
        executable = str(explicit) if explicit and explicit.is_file() else (shutil.which("pdftoppm.exe") or shutil.which("pdftoppm"))
        if not executable:
            return {"available": False, "engine": "", "pages": 0, "images": []}
        root = Path(tempfile.mkdtemp(prefix="clawos-pdf-render-")); prefix = root / "page"
        result = subprocess.run([executable, "-png", "-r", "144", str(path), str(prefix)], capture_output=True, text=True, timeout=120)
        images = sorted(str(item) for item in root.glob("page-*.png"))
        return {"available": result.returncode == 0 and bool(images), "engine": "pdftoppm", "pages": len(images), "images": images, "pdf": str(path), "stderr": result.stderr[-1000:]}

    @staticmethod
    def rendered_note_checks(source: Path, rendered_pdf: str | Path) -> list[dict[str, Any]]:
        """A visible note marker alone is not proof that its evidence text survived."""
        namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        notes = []
        with zipfile.ZipFile(source) as archive:
            for part, kind in (("word/footnotes.xml", "footnote"), ("word/endnotes.xml", "endnote")):
                if part not in archive.namelist():
                    continue
                for note in ET.fromstring(archive.read(part)).findall(f"{{{namespace}}}{kind}"):
                    if note.get(f"{{{namespace}}}type", "normal") != "normal" or int(note.get(f"{{{namespace}}}id", "-1")) < 1:
                        continue
                    value = "".join(node.text or "" for node in note.iter(f"{{{namespace}}}t"))
                    if value.strip():
                        notes.append(value)
        if not notes:
            return []
        error, text = "", ""
        try:
            import fitz
            with fitz.open(str(rendered_pdf)) as document:
                text = " ".join(page.get_text() for page in document)
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
        normalize = lambda value: re.sub(r"\s+", "", value).casefold()
        rendered = normalize(text)
        missing = [index for index, value in enumerate(notes, 1) if normalize(value) not in rendered]
        return [{"name": "rendered_note_content", "passed": not missing and not error, "severity": "high",
                 "detail": {"expected_notes": len(notes), "missing_note_indices": missing, "error": error}}]

    def _render_via_office(self, path: Path) -> dict[str, Any]:
        known = Path("C:/Program Files/LibreOffice/program/soffice.com")
        soffice = str(known) if known.is_file() else (shutil.which("soffice.exe") or shutil.which("soffice") or shutil.which("libreoffice"))
        if not soffice:
            return {"available": False, "engine": "", "pages": 0, "images": []}
        root = Path(tempfile.mkdtemp(prefix="clawos-office-render-"))
        # A running desktop LibreOffice instance can absorb headless conversion
        # commands and return 0 without producing a file. A per-run profile
        # keeps visual QA isolated and reproducible on Windows.
        profile = root / "profile"
        result = subprocess.run([
            soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless",
            "--convert-to", "pdf", "--outdir", str(root), str(path),
        ], capture_output=True, text=True, timeout=180)
        pdf = root / f"{path.stem}.pdf"
        return self._render_pdf(pdf) if result.returncode == 0 and pdf.is_file() else {"available": False, "engine": "libreoffice", "pages": 0, "images": [], "stderr": result.stderr[-1000:]}

    @staticmethod
    def _template_fidelity(path: Path, template: Path) -> dict[str, Any]:
        passed = template.is_file() and template.suffix.lower() == path.suffix.lower()
        detail: dict[str, Any] = {"template": str(template)}
        if passed and path.suffix.lower() == ".docx":
            from docx import Document
            a, b = Document(path), Document(template)
            source_styles = {p.style.name for p in b.paragraphs if p.style}
            output_styles = {p.style.name for p in a.paragraphs if p.style}
            ratio = len(source_styles & output_styles) / max(1, len(source_styles))
            detail["style_retention"] = ratio; passed = ratio >= 0.8
        return {"name": "template_fidelity", "passed": passed, "severity": "high", "detail": detail}

    @staticmethod
    def _build_workbook(rows: list[dict[str, Any]], output: Path, title: str, columns: list[str], metrics: list[str], summaries: dict[str, Any], grouped: dict[str, Any], group_by: str) -> None:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        wb = Workbook(); data = wb.active; data.title = "Data"
        data.append(columns)
        for row in rows: data.append([row.get(key) for key in columns])
        data.freeze_panes = "A2"; data.auto_filter.ref = data.dimensions
        table = Table(displayName="SourceData", ref=data.dimensions); table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        data.add_table(table)
        for cell in data[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="17365D")
        for index, key in enumerate(columns, 1): data.column_dimensions[data.cell(1,index).column_letter].width = min(32, max(12, len(key)+3))
        summary = wb.create_sheet("Summary"); summary.append([title]); summary["A1"].font = Font(size=20, bold=True, color="17365D")
        summary.append(["Metric", "Count", "Sum", "Mean", "Median", "Min", "Max", "Std Dev"])
        from openpyxl.utils import get_column_letter
        for row_no, key in enumerate(metrics, 3):
            col=get_column_letter(columns.index(key)+1); ref=f"'Data'!{col}2:{col}{len(rows)+1}"
            summary.append([key, f"=COUNT({ref})", f"=SUM({ref})", f"=AVERAGE({ref})", f"=MEDIAN({ref})", f"=MIN({ref})", f"=MAX({ref})", f"=STDEV.S({ref})"])
            for cell in summary[row_no][1:]: cell.number_format = "#,##0.00"
        summary.freeze_panes="A3"; summary.auto_filter.ref=f"A2:H{max(2,summary.max_row)}"
        for cell in summary[2]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="17365D")
        if grouped:
            group = wb.create_sheet("Group Analysis"); group.append([group_by,"Count",*metrics])
            group_col=get_column_letter(columns.index(group_by)+1); group_range=f"'Data'!{group_col}2:{group_col}{len(rows)+1}"
            for row_no, (name, item) in enumerate(grouped.items(), 2):
                values=[name, f'=COUNTIF({group_range},A{row_no})']
                for key in metrics:
                    metric_col=get_column_letter(columns.index(key)+1); metric_range=f"'Data'!{metric_col}2:{metric_col}{len(rows)+1}"
                    values.append(f'=SUMIF({group_range},A{row_no},{metric_range})')
                group.append(values)
            group.freeze_panes="A2"; group.auto_filter.ref=group.dimensions
            for cell in group[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="17365D")
            if group.max_row > 1 and metrics:
                chart=BarChart(); chart.title=f"{metrics[0]} by {group_by}"; chart.add_data(Reference(group,min_col=3,min_row=1,max_row=group.max_row),titles_from_data=True); chart.set_categories(Reference(group,min_col=1,min_row=2,max_row=group.max_row)); group.add_chart(chart,"F2")
        for ws in wb.worksheets:
            ws.sheet_view.showGridLines=False
            for row in ws.iter_rows():
                for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
        output.parent.mkdir(parents=True,exist_ok=True); wb.save(output)
