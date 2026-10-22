"""Reproduce Office edit/reopen/render evidence with installed local engines."""
import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from services.office_roundtrip_service import OfficeRoundtripService
from services.office_cross_application_service import OfficeCrossApplicationService


def evaluate(output: Path, cross_app: bool = False):
    output.mkdir(parents=True, exist_ok=True)
    service = OfficeRoundtripService()
    document = Document()
    document.add_heading("Office editing acceptance", 0)
    p = document.add_paragraph()
    p.add_run("Quarterly ").bold = True
    p.add_run("report").italic = True
    document.add_paragraph("Revenue is stable.")
    document.sections[0].header.paragraphs[0].text = "Quarterly report"
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"; table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Validated formulas"; table.cell(1, 1).text = "3"
    document.save(output / "source.docx")
    doc_edit = service.edit_docx(output / "source.docx", output / "edited.docx",
        replacements={"Quarterly report": "Updated report"},
        comments=[{"target": "Updated report", "text": "Cross-run edit verified"}],
        revisions=[{"old": "stable", "new": "growing"}])
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Data"
    for row in [["Input", "Value"], ["First", 10], ["Second", 20]]:
        sheet.append(row)
    sheet["B2"].number_format = "0.00"
    sheet["B4"] = "=SUM(B2:B3)"; sheet["B5"] = '=IF(B4>40,"passed","failed")'; sheet["B6"] = '=""'
    workbook.save(output / "source.xlsx")
    sheet_edit = service.edit_xlsx(output / "source.xlsx", output / "edited.xlsx", [{"sheet": "Data", "cell": "B2", "value": 30}])
    recalc = service.recalculate_xlsx(output / "edited.xlsx", output / "recalculated.xlsx")
    values = load_workbook(output / "recalculated.xlsx", data_only=True)
    formula_values = [values["Data"][address].value for address in ("B4", "B5", "B6")]
    values.close()
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.paragraphs[0]
    paragraph.add_run().text = "Quarterly "; paragraph.add_run().text = "report"
    deck.save(output / "source.pptx")
    slide_edit = service.edit_pptx(output / "source.pptx", output / "edited.pptx", {"Quarterly report": "Updated report"})
    renders = {}
    for name in ("edited.docx", "recalculated.xlsx", "edited.pptx"):
        target = output / (name + ".pdf")
        renders[name] = OfficeCrossApplicationService._libreoffice(output / name, target)
        if renders[name]["status"] == "completed":
            import fitz
            with fitz.open(target) as pdf:
                for index, page in enumerate(pdf):
                    page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3)).save(output / f"{name}-page-{index+1}.png")
    checks = {"docx_preserved": doc_edit["roundtrip_passed"], "xlsx_preserved": sheet_edit["roundtrip_passed"],
              "pptx_preserved": slide_edit["roundtrip_passed"], "formulas_complete": recalc["recalculated"],
              "formula_values_correct": formula_values[:2] == [50, "passed"] and formula_values[2] in ("", None),
              "all_formats_rendered": all(row["status"] == "completed" for row in renders.values())}
    report = {"checks": checks, "passed": all(checks.values()), "formula_values": formula_values,
              "recalculation": recalc, "renders": renders, "wps_parity_verified": False}
    if cross_app:
        report["cross_application"] = OfficeCrossApplicationService().evaluate(output / "edited.docx", output / "cross-app")
    (output / "acceptance.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--cross-app", action="store_true")
    args = parser.parse_args()
    result = evaluate(args.out.resolve(), args.cross_app)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["passed"] else 1)
