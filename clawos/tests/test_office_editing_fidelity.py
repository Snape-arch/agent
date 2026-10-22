import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches
from services.office_package_editing import parse_xml, serialize, S
from services.office_productivity_service import OfficeProductivityService
from services.office_roundtrip_service import OfficeRoundtripService, W


class OfficeEditingFidelityTests(unittest.TestCase):
    def test_cross_application_renderers_receive_distinct_snapshots(self):
        from services.office_cross_application_service import OfficeCrossApplicationService
        with tempfile.TemporaryDirectory() as td:
            source = Path(td)/"source.docx"
            doc = Document(); doc.add_paragraph("Original content"); doc.save(source)
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            observed = []
            def unavailable(path, target, *args):
                observed.append(path)
                self.assertNotEqual(path, source)
                self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), original_hash)
                return {"status": "unavailable"}
            with patch.object(OfficeCrossApplicationService, "_libreoffice", side_effect=unavailable), patch.object(OfficeCrossApplicationService, "_com_export", side_effect=unavailable):
                result = OfficeCrossApplicationService().evaluate(source, Path(td)/"render")
            self.assertEqual(len(set(observed)), 3)
            self.assertEqual(result["source_sha256"], original_hash)
            self.assertFalse(result["acceptance"]["passed"])

    def test_http_edit_forwards_review_and_recalculation_options(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        import clawos_server as server
        from unittest.mock import Mock
        application = FastAPI()
        for route in server.app.routes:
            if route.path == "/api/workspace/office/productivity":
                application.router.routes.append(route)
        service = Mock()
        service.edit.return_value = {"valid": True}
        with patch.object(server, "get_office_productivity_service", return_value=service), TestClient(application) as client:
            response = client.post("/api/workspace/office/productivity", json={"action": "edit", "file_path": "a.docx",
                "output_path": "b.docx", "confirmed": True, "comments": [{"target": "text", "text": "note"}],
                "revisions": [{"old": "old", "new": "new"}], "recalculate": True})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(service.edit.call_args.kwargs["recalculate"])
        self.assertEqual(len(service.edit.call_args.kwargs["comments"]), 1)
        self.assertEqual(len(service.edit.call_args.kwargs["revisions"]), 1)

    def test_partial_formula_caches_do_not_pass_recalculation(self):
        from unittest.mock import Mock
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td)/"in.xlsx", Path(td)/"out.xlsx"
            workbook = Workbook(); workbook.active["A1"] = "=1+1"; workbook.active["A2"] = "=2+2"; workbook.save(source)
            output.write_bytes(b"old output")
            def convert(args, **kwargs):
                target = Path(args[args.index("--outdir")+1]) / source.name
                parts = OfficeRoundtripService._read(source)
                root = parse_xml(parts["xl/worksheets/sheet1.xml"])
                root.find(f".//{{{S}}}v").text = "2"
                parts["xl/worksheets/sheet1.xml"] = serialize(root)
                OfficeRoundtripService._write(target, parts)
                return Mock(returncode=0)
            with patch("services.office_roundtrip_service.shutil.which", return_value="soffice"), patch("services.office_roundtrip_service.subprocess.run", side_effect=convert):
                with self.assertRaisesRegex(RuntimeError, "recalculation gate"):
                    OfficeRoundtripService().recalculate_xlsx(source, output)
            self.assertEqual(output.read_bytes(), b"old output")

    def test_comments_and_revisions_span_runs_without_duplicate_ids(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td)/"in.docx", Path(td)/"out.docx"
            doc = Document(); p = doc.add_paragraph()
            p.add_run("Annual ").bold = True; p.add_run("report")
            p = doc.add_paragraph(); p.add_run("Revenue ").italic = True; p.add_run("stable")
            doc.save(source)
            service = OfficeRoundtripService()
            service.edit_docx(source, output, comments=[{"target": "Annual report", "text": "Reviewed"}],
                revisions=[{"old": "Revenue stable", "new": "Revenue increasing"}])
            root = parse_xml(service._read(output)["word/document.xml"])
            self.assertEqual("".join(root.itertext()).count("Revenue increasing"), 1)
            self.assertEqual(len(list(root.iter(f"{{{W}}}commentRangeStart"))), 1)
            # A second revision uses fresh IDs and leaves previous review markup intact.
            service.edit_docx(output, output, revisions=[{"old": "Annual", "new": "Yearly"}])
            root = parse_xml(service._read(output)["word/document.xml"])
            ids = [node.get(f"{{{W}}}id") for node in root.iter() if node.tag in {f"{{{W}}}ins", f"{{{W}}}del"}]
            self.assertEqual(len(ids), len(set(ids)))

    def test_pdf_operations_reopen_and_invalid_batch_rolls_back(self):
        import fitz
        from services.office_pdf_editing import edit_pdf
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td)/"in.pdf", Path(td)/"out.pdf"
            with fitz.open() as doc:
                for number in range(3):
                    doc.new_page().insert_text((40, 40), f"Page {number+1}")
                doc.save(source)
            result = edit_pdf(source, output, [{"action": "reorder", "pages": [3, 1, 2]},
                {"action": "rotate", "pages": [1], "degrees": 90},
                {"action": "annotate", "pages": [2], "text": "Review note"},
                {"action": "delete", "pages": [3]}], lambda value: Path(value))
            self.assertEqual(result["pages"], 2)
            with fitz.open(output) as doc:
                self.assertIn("Page 3", doc[0].get_text())
                self.assertEqual(doc[0].rotation, 90)
                self.assertEqual(list(doc[1].annots())[0].info["content"], "Review note")
            previous = output.read_bytes()
            service = OfficeProductivityService(allowed_roots=[td])
            with self.assertRaises(ValueError):
                service.edit(source, output, updates=[{"action": "rotate", "pages": [1], "degrees": 90},
                    {"action": "delete", "pages": [1, 2, 3]}], confirmed=True)
            self.assertEqual(output.read_bytes(), previous)

    def test_docx_cross_runs_headers_and_same_path_source_hash(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "report.docx"
            doc = Document()
            p = doc.add_paragraph()
            p.add_run("Before annual ").bold = True
            p.add_run("report after").italic = True
            doc.sections[0].header.paragraphs[0].text = "annual report"
            doc.save(source)
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            result = OfficeRoundtripService().edit_docx(source, source, replacements={"annual report": "updated report"})
            reopened = Document(source)
            self.assertEqual(reopened.paragraphs[0].text, "Before updated report after")
            self.assertTrue(reopened.paragraphs[0].runs[0].bold)
            self.assertTrue(reopened.paragraphs[0].runs[1].italic)
            self.assertEqual(reopened.sections[0].header.paragraphs[0].text, "updated report")
            self.assertEqual(result["source_sha256"], before)
            self.assertNotEqual(result["source_sha256"], result["output_sha256"])
            self.assertEqual(result["replacement_counts"]["annual report"], 2)
            with zipfile.ZipFile(source) as archive:
                root = parse_xml(archive.read("word/document.xml"))
                self.assertIn("w14", root.nsmap)
                self.assertTrue(root.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"))

    def test_failed_replacement_and_broken_relation_do_not_replace_output(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td)/"in.docx", Path(td)/"out.docx"
            doc = Document(); doc.add_paragraph("first\tsecond"); doc.save(source)
            output.write_bytes(b"keep existing output")
            service = OfficeRoundtripService()
            with self.assertRaises(ValueError):
                service.edit_docx(source, output, replacements={"firstsecond": "bad"})
            self.assertEqual(output.read_bytes(), b"keep existing output")
            parts = service._read(source)
            parts.pop("word/styles.xml")
            service._write(source, parts)
            with self.assertRaisesRegex(ValueError, "preservation gate"):
                service.edit_docx(source, output, replacements={"first": "changed"})
            self.assertEqual(output.read_bytes(), b"keep existing output")

    def test_pptx_cross_runs_and_missing_target_are_verified(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td)/"in.pptx", Path(td)/"out.pptx"
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
            p = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame.paragraphs[0]
            p.add_run().text = "Quarterly "
            run = p.add_run(); run.text = "report"; run.font.bold = True
            deck.save(source)
            result = OfficeRoundtripService().edit_pptx(source, output, {"Quarterly report": "Updated results"})
            self.assertEqual(Presentation(output).slides[0].shapes[0].text, "Updated results")
            self.assertTrue(result["roundtrip_passed"])
            with self.assertRaises(ValueError):
                OfficeRoundtripService().edit_pptx(source, output, {"missing": "new"})

    def test_xlsx_real_relationship_styles_sorted_cells_and_cache_invalidation(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td)/"in.xlsx", Path(td)/"out.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "Data"
            ws["B3"] = 5; ws["B3"].font = Font(bold=True); ws["B3"].number_format = "0.00%"
            ws["B3"].fill = PatternFill("solid", fgColor="00FF00")
            ws.merge_cells("D1:E1"); ws["D1"] = "Merged"
            summary = wb.create_sheet("Summary"); summary["A1"] = "=Data!B3*2"
            wb.save(source)
            service = OfficeRoundtripService(); parts = service._read(source)
            # Workbook relationship points to a non-sequential worksheet part.
            parts["xl/worksheets/custom.xml"] = parts.pop("xl/worksheets/sheet1.xml")
            parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(b"sheet1.xml", b"custom.xml")
            parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(b"sheet1.xml", b"custom.xml")
            root = parse_xml(parts["xl/worksheets/sheet2.xml"])
            root.find(f".//{{{S}}}v").text = "10"
            parts["xl/worksheets/sheet2.xml"] = serialize(root)
            service._write(source, parts)
            result = service.edit_xlsx(source, output, [{"sheet": "Data", "cell": "B3", "value": .75},
                {"sheet": "Data", "cell": "A1", "value": "=literal", "type": "text"}])
            reopened = load_workbook(output)
            self.assertEqual(reopened["Data"]["B3"].value, .75)
            self.assertTrue(reopened["Data"]["B3"].font.bold)
            self.assertEqual(reopened["Data"]["B3"].number_format, "0.00%")
            self.assertEqual(reopened["Data"]["A1"].data_type, "s")
            self.assertEqual(reopened["Summary"]["A1"].value, "=Data!B3*2")
            reopened.close()
            self.assertEqual(service._formula_counts(output), (1, 0))
            self.assertTrue(result["roundtrip_passed"])
            original_output = output.read_bytes()
            for update in ({"sheet": "typo", "cell": "A1", "value": 1},
                           {"sheet": "Data", "cell": "E1", "value": 2},
                           {"sheet": "Data", "cell": "XFE1", "value": 2},
                           {"sheet": "Data", "cell": "A0", "value": 2},
                           {"sheet": "Data", "cell": "A1", "value": float("nan")}):
                with self.assertRaises(ValueError):
                    service.edit_xlsx(source, output, [update])
                self.assertEqual(output.read_bytes(), original_output)

    def test_operation_recalculation_failure_preserves_source_and_destination(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td)/"in.xlsx", Path(td)/"out.xlsx"
            wb = Workbook(); wb.active["A1"] = 2; wb.save(source)
            original = source.read_bytes(); output.write_bytes(b"old destination")
            service = OfficeProductivityService(allowed_roots=[td])
            with patch.object(service.roundtrip, "recalculate_xlsx", side_effect=RuntimeError("failed")):
                with self.assertRaises(RuntimeError):
                    service.edit(source, output, updates=[{"sheet": "Sheet", "cell": "A1", "value": 7}], recalculate=True, confirmed=True)
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(output.read_bytes(), b"old destination")


if __name__ == "__main__":
    unittest.main()
