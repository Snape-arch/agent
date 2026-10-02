import asyncio
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import httpx
from fastapi import FastAPI, HTTPException
from memoryos.memory_store import MemoryStore
from services.academy_documents import AcademyDocuments,documents_router
from services.collaboration import Collaboration,AuthorizedDocuments,collaboration_router
from services.data_lab import DataLab,process,profile,chart,parse_upload,safe_cell
from services.meeting_workbench import MeetingWorkbench
from workspace_services import WorkspaceServices
from security.policy import ToolSecurityPolicy


class WorkbenchTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.store=MemoryStore(str(Path(self.tmp.name)/'db.sqlite'))
        self.docs=AcademyDocuments(self.store);self.access=Collaboration(self.store)
        self.doc=self.docs.save(dict(title='Private evidence',blocks=[dict(id='a',type='paragraph',text='source A')]))
        self.data=DataLab(self.store)
        self.meetings=MeetingWorkbench(WorkspaceServices(self.store,ToolSecurityPolicy([self.tmp.name])))

    def member(self,role='editor'):
        invite=self.access.invite(self.doc['id'],'local-user',role)
        token,user=self.access.join(invite['code'],'Collaborator')
        return token,user,AuthorizedDocuments(self.docs,self.access,user['id'])

    def test_invite_single_use_and_stored_tokens_hashed(self):
        invite=self.access.invite(self.doc['id'],'local-user','viewer')
        def join():
            try:return self.access.join(invite['code'],'Same invitation')[1]
            except HTTPException as e:return e.status_code
        with ThreadPoolExecutor(2) as pool:results=list(pool.map(lambda _:join(),range(2)))
        self.assertEqual(sum(isinstance(r,dict) for r in results),1)
        self.assertIn(403,results)
        with self.store._conn() as c:
            self.assertNotEqual(c.execute('SELECT hash FROM collaboration_invites').fetchone()['hash'],invite['code'])

    def test_reader_cannot_edit_or_claim_ownership_and_revocation_immediate(self):
        _,u,s=self.member('viewer');doc=s.get(self.doc['id'],'local-user')
        with self.assertRaises(HTTPException) as error:s.save(doc,self.doc['id'],'local-user')
        self.assertEqual(error.exception.status_code,403)
        other=self.docs.save(dict(title='Other',blocks=[dict(id='b',type='paragraph',text='hidden')]))
        with self.assertRaises(HTTPException):s.get(other['id'])
        self.assertEqual(len(s.index('local-user')['documents']),1)
        self.access.revoke(self.doc['id'],'local-user',u['id'])
        with self.assertRaises(HTTPException):s.history(self.doc['id'])

    def test_editor_save_is_cas_authorized_and_audited(self):
        _,u,s=self.member();doc=s.get(self.doc['id']);doc['blocks'][0]['text']='reviewed'
        updated=s.save(doc,self.doc['id']);self.assertEqual(updated['version'],2)
        from services.academy_documents import DocumentConflict
        with self.assertRaises(DocumentConflict):s.save(doc,self.doc['id'])
        state=self.access.state(self.doc['id'],'local-user')
        self.assertTrue(any(a['action']=='edit' and a['actor']==u['id'] for a in state['audit']))
        with self.assertRaises(HTTPException):s.save(doc)

    def test_comments_bound_to_actual_block_version_and_role(self):
        _,u,_=self.member('commenter')
        cid=self.access.comment(self.doc['id'],u['id'],dict(text='check evidence',block_id='a',version=1))['id']
        self.assertEqual(self.access.state(self.doc['id'],u['id'])['comments'][0]['id'],cid)
        with self.assertRaises(HTTPException):self.access.comment(self.doc['id'],u['id'],dict(text='bad anchor',block_id='foreign',version=1))
        with self.assertRaises(HTTPException):self.access.comment(self.doc['id'],u['id'],dict(text='stale',block_id='a',version=2))
        _,v,_=self.member('viewer')
        with self.assertRaises(HTTPException):self.access.comment(self.doc['id'],v['id'],dict(text='no',block_id='a',version=1))

    def test_remote_document_api_ignores_client_identity(self):
        async def run():
            app=FastAPI();app.include_router(documents_router(lambda:self.docs,None,lambda:self.access));app.include_router(collaboration_router(lambda:self.access))
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app,client=('203.0.113.5',3210)),base_url='http://test') as client:
                base='/api/academy/learning/documents'
                self.assertEqual((await client.get(base+'?learner_id=local-user')).status_code,401)
                invite=self.access.invite(self.doc['id'],'local-user','viewer')
                joined=await client.post('/api/collaboration/join',json=dict(code=invite['code'],name='Reader'))
                self.assertIn('HttpOnly',joined.headers['set-cookie']);self.assertIn('SameSite=strict',joined.headers['set-cookie'])
                doc=(await client.get(base+'/'+self.doc['id'])).json()
                self.assertEqual((await client.put(base+'/'+self.doc['id'],json={**doc,'learner_id':'local-user'})).status_code,403)
                self.assertEqual((await client.get(base+'/'+self.doc['id']+'/export')).status_code,200)
                await client.post('/api/collaboration/logout')
                self.assertEqual((await client.get(base+'/'+self.doc['id'])).status_code,401)
                invite=self.access.invite(self.doc['id'],'local-user','viewer')
                await client.post('/api/collaboration/join',json=dict(code=invite['code'],name='API reader'))
                token=client.cookies.get('clawos_session');client.cookies.clear()
                headers={'Authorization':'Bearer '+token}
                self.assertEqual((await client.get(base+'/'+self.doc['id'],headers=headers)).status_code,200)
                await client.post('/api/collaboration/logout',headers=headers)
                self.assertEqual((await client.get(base+'/'+self.doc['id'],headers=headers)).status_code,401)
        asyncio.run(run())

    def test_csv_encoding_headers_and_xlsx_formulas_are_not_executed(self):
        cols,rows,_=parse_upload('a.csv','团队,值\n甲,0\n乙,2\n'.encode('gb18030'))
        self.assertEqual(cols,['团队','值']);self.assertEqual({p['x']:p['y'] for p in chart(cols,rows,dict(kind='bar',x='团队',y='值'))['points']}['甲'],0)
        with self.assertRaises(ValueError):parse_upload('bad.csv',b'a,a\n1,2')
        with self.assertRaises(ValueError):parse_upload('bad.csv',b'a,b\n1,2,3')
        from openpyxl import Workbook
        wb=Workbook();ws=wb.active;ws.append(['name','value']);ws.append(['x','=1+1']);buf=io.BytesIO();wb.save(buf)
        _,r,warnings=parse_upload('a.xlsx',buf.getvalue());self.assertIsNone(r[0]['value']);self.assertTrue(warnings)
        self.assertEqual(safe_cell('  =cmd'),"'  =cmd")

    def test_pipeline_fill_dedup_derive_group_and_source_replay(self):
        cols=['group','x','y'];rows=[dict(group='A',x=2,y=4),dict(group='A',x=2,y=4),dict(group='B',x=None,y=6)]
        d=self.data.create('example',cols,rows)
        recipe=[dict(op='deduplicate'),dict(op='fill',column='x',strategy='mean'),dict(op='derive',left='y',right='x',name='ratio',method='divide'),dict(op='group',group='group',metric='ratio',aggregate='mean')]
        r=self.data.run(d['id'],1,recipe);self.assertEqual([x['value'] for x in r['rows']],[2,3])
        self.assertEqual(len(r['steps']),4)
        restored=self.data.run(d['id'],2,[]);self.assertEqual(restored['profile']['rows'],3);self.assertEqual(restored['profile']['duplicates'],1)
        self.assertEqual(d['source_hash'],restored['source_hash'])

    def test_invalid_pipeline_rolls_back_and_stale_version_rejected(self):
        d=self.data.create('bad',['a'],[dict(a='not numeric')])
        with self.assertRaises(ValueError):self.data.run(d['id'],1,[dict(op='numeric',column='a')])
        self.assertEqual(self.data.get(d['id'])['version'],1)
        with self.assertRaises(HTTPException):self.data.run(d['id'],999,[])
        with self.store._conn() as c:self.assertEqual(c.execute('SELECT count(*) FROM data_lab_runs').fetchone()[0],0)

    def test_charts_statistics_use_all_rows_not_preview_and_exclude_missing(self):
        rows=[dict(x=i,y=2*i) for i in range(250)]+[dict(x=None,y=4)]
        d=self.data.create('large',['x','y'],rows)
        self.assertEqual(len(d['rows']),100);self.assertEqual(d['profile']['rows'],251)
        full=self.data.get(d['id'],True)
        c=chart(full['columns'],full['rows'],dict(kind='scatter',x='x',y='y'))
        self.assertAlmostEqual(c['correlation'],1);self.assertEqual(c['omitted'],1);self.assertEqual(len(c['points']),250)
        hist=chart(full['columns'],full['rows'],dict(kind='histogram',x='x'))
        self.assertEqual(sum(p['y'] for p in hist['points']),250)

    def test_http_identity_scope_csrf_and_export_full_data(self):
        from unittest.mock import patch
        from clawos_server import workbench_access
        from services.data_lab import data_lab_router
        async def run():
            app=FastAPI();app.middleware('http')(workbench_access)
            app.include_router(documents_router(lambda:self.docs,None,lambda:self.access))
            app.include_router(data_lab_router(lambda:self.data,lambda:self.access))
            token,u,_=self.member('viewer')
            with patch('clawos_server._collaboration',return_value=self.access):
                async with httpx.AsyncClient(transport=httpx.ASGITransport(app,client=('203.0.113.1',1000)),base_url='http://test') as remote:
                    self.assertEqual((await remote.get('/api/data-lab')).status_code,401)
                    self.assertEqual((await remote.get('/api/data-lab',headers={'Authorization':'Bearer '+token})).status_code,403)
                    r=await remote.get('/api/academy/learning/documents/'+self.doc['id'],headers={'Authorization':'Bearer '+token})
                    self.assertEqual(r.status_code,200);self.assertEqual(r.json()['role'],'viewer')
                async with httpx.AsyncClient(transport=httpx.ASGITransport(app,client=('127.0.0.1',1000)),base_url='http://test') as owner:
                    self.assertEqual((await owner.get('/api/data-lab',headers={'X-Forwarded-For':'203.0.113.2'})).status_code,401)
                    self.assertEqual((await owner.post('/api/data-lab',json={},headers={'Origin':'https://foreign.example'})).status_code,403)
                    self.assertEqual((await owner.post('/api/data-lab',json={},headers={'Sec-Fetch-Site':'cross-site'})).status_code,403)
                    d=self.data.create('all rows',['n','text'],[dict(n=i,text='=cmd' if i==0 else 'safe') for i in range(250)])
                    csv=await owner.get('/api/data-lab/'+d['id']+'/export')
                    self.assertEqual(len(csv.text.splitlines()),251);self.assertIn("'=cmd",csv.text)
                    xlsx=await owner.get('/api/data-lab/'+d['id']+'/export?format=xlsx')
                    from openpyxl import load_workbook
                    wb=load_workbook(io.BytesIO(xlsx.content),data_only=False)
                    self.assertEqual(wb['Data'].max_row,251);self.assertEqual(wb['Data']['B2'].data_type,'s');self.assertEqual(wb['Recipe']['B1'].value,d['source_hash'])
        asyncio.run(run())

    def test_blocked_document_write_does_not_stop_other_http_requests(self):
        import threading
        from unittest.mock import patch
        started,release=threading.Event(),threading.Event();original=self.docs.save
        def blocked_save(*args,**kwargs):
            started.set()
            if not release.wait(3):raise RuntimeError('test barrier not released')
            return original(*args,**kwargs)
        async def run():
            app=FastAPI();app.include_router(documents_router(lambda:self.docs,None,lambda:self.access))
            @app.get('/ping')
            async def ping():return {'ready':True}
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app,client=('127.0.0.1',1234)),base_url='http://test') as client:
                with patch.object(self.docs,'save',side_effect=blocked_save):
                    pending=asyncio.create_task(client.put('/api/academy/learning/documents/'+self.doc['id'],json=self.doc))
                    try:
                        for _ in range(200):
                            if started.is_set():break
                            await asyncio.sleep(.005)
                        self.assertTrue(started.is_set());self.assertFalse(pending.done())
                        response=await asyncio.wait_for(client.get('/ping'),1)
                        self.assertEqual(response.status_code,200);self.assertFalse(pending.done())
                    finally:release.set()
                    self.assertEqual((await pending).status_code,200)
        asyncio.run(run())

    def test_parallel_edits_and_pipeline_have_one_version_winner(self):
        from services.academy_documents import DocumentConflict
        _,_,s=self.member();doc=s.get(self.doc['id'])
        def save(i):
            payload=json.loads(json.dumps(doc));payload['blocks'][0]['text']=str(i)
            try:return s.save(payload,doc['id'])['version']
            except DocumentConflict:return 'conflict'
        with ThreadPoolExecutor(4) as pool:r=list(pool.map(save,range(4)))
        self.assertEqual(r.count(2),1);self.assertEqual(r.count('conflict'),3)
        d=self.data.create('version',['x'],[dict(x=1)])
        def pipeline(i):
            try:return self.data.run(d['id'],1,[])['version']
            except HTTPException as e:return e.status_code
        with ThreadPoolExecutor(4) as pool:r=list(pool.map(pipeline,range(4)))
        self.assertEqual(r.count(2),1);self.assertEqual(r.count(409),3)

    def test_meeting_document_snapshot_preserves_source_and_is_idempotent(self):
        m=self.meetings.draft('Snapshot','决定：保留证据\n行动项：Alice | 复核 | 2026-10-10')
        link=self.meetings.document(m['id']);again=self.meetings.document(m['id'])
        self.assertEqual(link,again);doc=self.docs.get(link['document_id'])
        self.assertIn(m['review']['source_hash'],doc['context'])
        self.assertIn('待复核草稿',doc['blocks'][0]['text'])
        changed=self.meetings.draft(m['title'],m['transcript']+'\n决定：新版本',m['id'],m['review']['version'])
        next_link=self.meetings.document(m['id']);self.assertNotEqual(link['document_id'],next_link['document_id'])
        self.assertNotIn('新版本',json.dumps(self.docs.get(link['document_id']),ensure_ascii=False))

    def test_meeting_publish_and_document_link_roll_back_as_one_unit(self):
        from unittest.mock import patch
        m=self.meetings.draft('Atomic source','决定：保留原文')
        with patch.object(self.meetings,'event',side_effect=RuntimeError('interrupted audit')):
            with self.assertRaises(RuntimeError):self.meetings.review(m['id'],m['review']['version'],'逐项核对全部原文')
            with self.assertRaises(RuntimeError):self.meetings.document(m['id'])
        self.assertTrue(self.meetings.get(m['id'])['review']['review_required'])
        with self.store._conn() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM knowledge_records_v2').fetchone()[0],0)
            self.assertEqual(c.execute('SELECT count(*) FROM meeting_document_links').fetchone()[0],0)
            self.assertEqual(c.execute('SELECT count(*) FROM academy_documents').fetchone()[0],1)

    def test_meeting_feishu_draft_requires_review_and_retains_source(self):
        from services.meeting_workbench import meeting_router
        from services.feishu import FeishuConfig,FeishuConnectorService
        feishu=FeishuConnectorService(self.store,None,FeishuConfig())
        async def run():
            app=FastAPI();app.include_router(meeting_router(lambda:self.meetings,lambda:self.access,lambda:feishu))
            m=self.meetings.draft('Approved snapshot','决定：保留依据')
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app,client=('127.0.0.1',1000)),base_url='http://test') as client:
                endpoint=f'/api/meeting-workbench/{m["id"]}/feishu-draft'
                self.assertEqual((await client.post(endpoint)).status_code,409)
                reviewed=self.meetings.review(m['id'],m['review']['version'],'已核对全部原文')
                r=await client.post(endpoint);self.assertEqual(r.status_code,200)
                self.assertIn(reviewed['review']['source_hash'],r.json()['source_message_ids'][0])
                self.assertEqual(feishu.list_approvals(),[])
        asyncio.run(run())

    def test_meeting_review_publish_then_source_change_expires_knowledge(self):
        transcript='Alice: 决定：发布验证版\nBob: 行动项：Bob | 完成回归 | 2026-10-10\n风险：资料不完整'
        m=self.meetings.draft('Review',transcript)
        self.assertTrue(m['review']['review_required']);self.assertIsNone(m['review']['knowledge_id'])
        reviewed=self.meetings.review(m['id'],m['review']['version'],'逐项核对原文，已确认负责人与期限')
        self.assertFalse(reviewed['review']['review_required']);kid=reviewed['review']['knowledge_id']
        again=self.meetings.review(m['id'],reviewed['review']['version'],'重复核对不重复入库')
        self.assertEqual(again['review']['knowledge_id'],kid)
        changed=self.meetings.draft('Review',transcript+'\n决定：补充证据后再发布',m['id'],again['review']['version'])
        self.assertTrue(changed['review']['review_required'])
        with self.store._conn() as c:self.assertEqual(c.execute('SELECT status FROM knowledge_records_v2 WHERE id=?',(kid,)).fetchone()['status'],'expired')

    def test_action_progress_versions_and_reprocessing_preserves_existing_status(self):
        m=self.meetings.draft('Actions','待办：Alice | 测试 | 2026-01-01');a=m['actions'][0]
        updated=self.meetings.action(m['id'],a['id'],m['review']['version'],dict(owner='Alice',due_date='2026-01-01',status='done'))
        self.assertEqual(updated['metrics']['done'],1);self.assertEqual(updated['metrics']['overdue'],0)
        with self.assertRaises(HTTPException):self.meetings.action(m['id'],a['id'],m['review']['version'],dict(owner='Alice',due_date='',status='open'))
        updated=self.meetings.draft('Actions',m['transcript'],m['id'],updated['review']['version'])
        self.assertEqual(updated['actions'][0]['status'],'done')

    def test_audio_repeated_title_does_not_overwrite_runs_or_publish_draft(self):
        import wave
        audio=Path(self.tmp.name)/'a.wav'
        with wave.open(str(audio),'wb') as w:w.setnchannels(1);w.setsampwidth(2);w.setframerate(16000);w.writeframes(b'\0\0'*1600)
        from services.meeting_audio_service import MeetingAudioService
        service=MeetingAudioService(self.store,Path(self.tmp.name)/'artifacts')
        a=service.process(audio_path=audio,title='same',transcript='Alice: 已收到')
        b=service.process(audio_path=audio,title='same',transcript='Alice: 第二次记录')
        self.assertNotEqual(a['artifacts']['meeting_card']['path'],b['artifacts']['meeting_card']['path'])
        self.assertNotEqual(a['run_record']['run_id'],b['run_record']['run_id'])
        self.assertIsNone(a['knowledge_item']);self.assertEqual(a['run_record']['status'],'needs_review')
        self.assertEqual(a['run_record']['workflow']['state_machine']['state'],'waiting_gate')
        self.assertTrue(any(g['name']=='human_review' and g['status']=='pending' for g in a['run_record']['workflow']['state_machine']['gates']))


if __name__=='__main__':unittest.main()
