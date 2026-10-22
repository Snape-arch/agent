import asyncio
import copy
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import AsyncMock, patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from memoryos.memory_store import MemoryStore
from services.academy_materials import AcademyMaterials
from services.academy_material_api import materials_router

TEXT='# 检索学习\n\n检索系统只在证据充分时回答。没有证据时必须拒答，不能猜测。\n\n# 参数\n\n器材借用期限为 7 天。活动人数最多 12 人。'
ANSWER='检索系统只有在证据支持答案时才作答。缺少资料时应拒绝猜测；例如没有费用规定，就不能从开放时间推测费用。'


def judge(score=2):
    return dict(dimensions={key:dict(score=score,reason='原文要求有证据，学生解释了缺失证据的反例。',student_quote='缺少资料时应拒绝猜测',source_quote='没有证据时必须拒答') for key in ('mechanism','conditions','counterexample')},contradiction=False,feedback='继续检查反例是否真正违反证据条件。')


class MaterialTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.store=MemoryStore(str(Path(self.tmp.name)/'db.sqlite'))
        self.service=AcademyMaterials(self.store,Path(self.tmp.name)/'materials')
        self.state=self.service.upload('教材.md',TEXT.encode())
        self.mid=self.state['document']['id'];self.cid=self.state['document']['payload']['passages'][0]['id']

    def test_version_and_scope_preserve_original(self):
        self.assertEqual(self.state['document']['status'],'ready')
        self.assertEqual(self.service.upload('教材.md',TEXT.encode())['document']['id'],self.mid)
        new=self.service.upload('教材.md',(TEXT+'\n新增条件').encode())
        self.assertNotEqual(new['document']['id'],self.mid)
        self.assertEqual(len(self.service.state()['materials']),2)
        with self.assertRaises(ValueError):self.service.document(self.mid,'other')
        self.assertEqual((self.service.root/self.mid/'教材.md').read_text(encoding='utf-8'),TEXT)

    def test_parser_failure_retry_and_page_provenance(self):
        with patch('services.academy_materials.extract_text',side_effect=RuntimeError('OCR offline')):
            bad=self.service.upload('扫描.pdf',b'fixture')
        self.assertEqual(bad['document']['status'],'failed')
        with patch('services.academy_materials.extract_text',return_value=('第一页正文\f第二页正文',{'extractor':'fixture','page_count':2,'ocr_required':True})):
            ready=self.service.parse(bad['document']['id'])
        self.assertEqual([p['page'] for p in ready['document']['payload']['passages']],[1,2])
        self.assertTrue(ready['document']['payload']['warnings'])

    def test_actual_docx_and_pdf_parser(self):
        from docx import Document
        import fitz
        doc=Document();doc.add_heading('Learning',0);doc.add_paragraph('Evidence is required. Missing evidence requires refusal.')
        buf=io.BytesIO();doc.save(buf)
        self.assertEqual(self.service.upload('book.docx',buf.getvalue())['document']['status'],'ready')
        pdf=fitz.open();p=pdf.new_page();p.insert_text((72,72),'Evidence is required before answering.');p=pdf.new_page();p.insert_text((72,72),'Missing evidence requires refusal.')
        parsed=self.service.upload('paper.pdf',pdf.tobytes());pdf.close()
        self.assertEqual(parsed['document']['status'],'ready')
        self.assertEqual({p['page'] for p in parsed['document']['payload']['passages']},{1,2})

    def test_guide_rejects_invented_quote(self):
        value=dict(summary='检索回答需要证据。',quote='没有证据时必须拒答',question='解释为什么没有证据时必须拒答，并举例。',prerequisites=['证据的含义'],steps=['先读原文再解释条件'])
        saved=self.service.save_guide(self.mid,self.cid,value)
        self.assertEqual(saved['data']['citation'],self.cid)
        with self.assertRaises(ValueError):self.service.save_guide(self.mid,self.cid,dict(value,quote='有证据就必须猜测'))
        self.assertEqual(self.service.lesson(self.mid,self.cid)['question'],value['question'])

    def test_diagnosis_and_tutor_context_include_remediation(self):
        result=self.service.diagnose(self.mid,{'source':0,'missing':0,'test':0},'理解证据')
        self.assertEqual(result['data']['result']['passed'],1)
        prompt=self.service.teaching_prompt(self.mid,self.cid,'为什么不能根据经验猜测？')
        self.assertIn('补学证据支持',prompt[1]['content'])
        with self.assertRaises(ValueError):self.service.teaching_prompt(self.mid,'fake','请解释资料内容')

    def test_semantic_agreement_contradiction_and_disagreement(self):
        result=self.service.save_assessment(self.mid,self.cid,ANSWER,[judge(),judge()])
        self.assertEqual(result['data']['status'],'passed')
        contrad=judge();contrad['contradiction']=True
        self.assertEqual(self.service.save_assessment(self.mid,self.cid,ANSWER,[contrad,contrad])['data']['status'],'needs_revision')
        self.assertEqual(self.service.save_assessment(self.mid,self.cid,ANSWER,[judge(),judge(1)])['data']['status'],'needs_review')
        self.assertEqual(self.service.state(self.mid)['events'][0]['kind'],'assessment')

    def test_semantic_evidence_must_exist_in_both_sources(self):
        for key,value in [('student_quote','学生根本没有写这句话'),('source_quote','这不是原文')]:
            bad=judge();bad['dimensions']['mechanism'][key]=value
            with self.subTest(key=key),self.assertRaises(ValueError):self.service.save_assessment(self.mid,self.cid,ANSWER,[bad,judge()])
        bad=judge();bad['dimensions']['mechanism']['score']=True
        with self.assertRaises(ValueError):self.service.save_assessment(self.mid,self.cid,ANSWER,[bad,judge()])

    def test_actual_project_task_and_export_reproduce(self):
        cases=[dict(query='器材借用期限',citation=self.cid,quote='7 天'),dict(query='活动人数最多',citation=self.cid,quote='12 人'),dict(query='火星天气预测',citation=None,quote='')]
        cases[0]['quote']='期限为 7 天';cases[1]['quote']='最多 12 人'
        result=self.service.project(self.mid,dict(chunk_size=160,top_k=2,min_score=.3,current_only=True),cases)['data']
        self.assertEqual(result['report']['passed'],3)
        task=self.store.get_task(result['task_id'])
        self.assertEqual(task['task']['status'],'succeeded')
        self.assertIn(result['artifact_digest'],json.dumps(task))
        archive=zipfile.ZipFile(io.BytesIO(self.service.export(self.mid,result['task_id'])))
        self.assertEqual(set(archive.namelist()),{'project.py','input.json','report.json','README.md'})
        root=Path(self.tmp.name)/'export';archive.extractall(root)
        import subprocess
        proc=subprocess.run([sys.executable,'project.py'],cwd=root,capture_output=True,timeout=10)
        self.assertEqual(proc.returncode,0,proc.stderr)
        self.assertEqual(json.loads((root/'report.json').read_text(encoding='utf-8')),result['report'])
        with self.assertRaises(ValueError):self.service.export(self.mid,result['task_id'],'other')
        (self.service.root/self.mid/f"run-{result['task_id']}"/'report.json').write_text('{}',encoding='utf-8')
        with self.assertRaises(ValueError):self.service.export(self.mid,result['task_id'])

    def test_project_failure_does_not_create_success(self):
        cases=[dict(query='器材借用期限',citation=self.cid,quote='期限为 7 天'),dict(query='活动人数最多',citation=self.cid,quote='最多 12 人'),dict(query='火星天气预测',citation=None,quote='')]
        with patch('services.coding_workflow_service.CodingWorkflowService._run_command',return_value=dict(exit_code=124,stderr='timeout')):
            with self.assertRaises(ValueError):self.service.project(self.mid,dict(chunk_size=160,top_k=2,min_score=.3,current_only=True),cases)
        self.assertFalse(any(e['kind']=='project' for e in self.service.state(self.mid)['events']))
        self.assertEqual(self.store.list_tasks()[0]['status'],'failed')

    def test_http_upload_assess_and_provider_failure(self):
        from fastapi import FastAPI
        import httpx
        model=AsyncMock(return_value={'choices':[{'message':{'content':json.dumps(judge(),ensure_ascii=False)}}]})
        app=FastAPI();app.include_router(materials_router(lambda:self.service,model))
        async def run():
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as c:
                uploaded=await c.post('/api/academy/learning/materials',files={'file':('fresh.md',TEXT.encode(),'text/plain')})
                self.assertEqual(uploaded.status_code,200)
                path='/api/academy/learning/materials/'+self.mid
                response=await c.post(path+'/assess',json=dict(citation=self.cid,answer=ANSWER))
                self.assertEqual(response.status_code,200,response.text)
                self.assertEqual(response.json()['data']['status'],'passed')
                self.assertEqual(model.call_count,2);self.assertEqual(model.call_args.args[1],[])
                before=len([e for e in self.service.state(self.mid)['events'] if e['kind']!='model_run'])
                model.side_effect=RuntimeError('offline')
                response=await c.post(path+'/guide',json=dict(citation=self.cid))
                self.assertEqual(response.status_code,502)
                self.assertEqual(before,len([e for e in self.service.state(self.mid)['events'] if e['kind']!='model_run']))
                self.assertEqual(self.service.state(self.mid)['events'][0]['data']['status'],'failed')
                denied=await c.get(path+'?learner_id=other');self.assertEqual(denied.status_code,404)
        asyncio.run(run())


if __name__=='__main__':unittest.main()
