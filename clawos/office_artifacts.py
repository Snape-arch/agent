"""Office helpers with professional engines and self-contained fallbacks.

When openpyxl or python-pptx are installed, ClawOS uses them for richer
offline artifacts. The stdlib OpenXML path remains as the no-dependency
fallback for fresh Python installs.
"""

from __future__ import annotations

import html
import os
import re
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
AML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"

ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)
ET.register_namespace("p", PML_NS)
ET.register_namespace("a", AML_NS)
ET.register_namespace("w", WML_NS)
ET.register_namespace("m", OMML_NS)


def create_docx(path, title="", sections=None, metadata=None):
    """Create a structured DOCX with a dependency-free OpenXML fallback."""
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        section = document.sections[0]
        # standard_business_brief preset: explicit Letter geometry and furniture.
        section.page_width = Inches(8.5); section.page_height = Inches(11)
        section.top_margin = Inches(1); section.bottom_margin = Inches(1)
        section.left_margin = Inches(1); section.right_margin = Inches(1)
        section.header_distance = Inches(0.492); section.footer_distance = Inches(0.492)
        styles = document.styles
        styles["Normal"].font.name = "Calibri"
        styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles["Normal"].font.size = Pt(11)
        styles["Normal"].paragraph_format.line_spacing = 1.10
        styles["Normal"].paragraph_format.space_after = Pt(6)
        style_tokens = {
            "Title": (24, "17365D", 0, 14),
            "Heading 1": (16, "2E74B5", 16, 8),
            "Heading 2": (13, "2E74B5", 12, 6),
            "Heading 3": (12, "1F4D78", 8, 4),
        }
        for name, (size, color, before, after) in style_tokens.items():
            style = styles[name]; style.font.name = "Calibri"; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color); style.font.bold = True
            style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri"); style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        title_p = document.add_paragraph(style="Title"); title_p.add_run(str(title or target.stem)); title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_p.paragraph_format.keep_with_next = True
        subtitle = str((metadata or {}).get("subtitle") or "")
        if subtitle:
            p=document.add_paragraph(subtitle); p.style=styles["Subtitle"]
        for item in sections or []:
            if item.get("page_break"):
                document.add_page_break()
            heading = str(item.get("heading") or "").strip()
            if heading:
                heading_p = document.add_heading(heading, level=max(1, min(int(item.get("level") or 1), 3)))
                heading_p.paragraph_format.keep_with_next = True
            for paragraph in item.get("paragraphs") or []:
                value = str(paragraph.get("text") or "") if isinstance(paragraph, dict) else str(paragraph)
                _add_docx_paragraph(document, value)
            bullets = item.get("bullets") or []
            for bullet in bullets:
                document.add_paragraph(str(bullet), style="List Bullet")
            table_spec = item.get("table") or []
            table = (table_spec.get("rows") or []) if isinstance(table_spec, dict) else table_spec
            if table:
                width = max(len(row) for row in table)
                word_table = document.add_table(rows=0, cols=width)
                word_table.style = "Table Grid"; word_table.autofit = False
                usable_inches = 6.5
                requested_widths = list(table_spec.get("column_widths") or []) if isinstance(table_spec, dict) else []
                if len(requested_widths) != width or sum(float(value or 0) for value in requested_widths) <= 0:
                    lengths = [max(3, max((len(str(row[index])) if index < len(row) else 0) for row in table)) for index in range(width)]
                    total = sum(lengths)
                    requested_widths = [usable_inches * value / total for value in lengths]
                else:
                    requested_widths = [usable_inches * float(value) / sum(float(item) for item in requested_widths) for value in requested_widths]
                for row_index, row in enumerate(table):
                    cells = word_table.add_row().cells
                    header_rows = int(table_spec.get("header_rows", 1)) if isinstance(table_spec, dict) else 1
                    if row_index < header_rows:
                        trPr = cells[0]._tc.getparent().get_or_add_trPr(); marker=OxmlElement("w:tblHeader"); marker.set(qn("w:val"), "true"); trPr.append(marker)
                    for col_index, value in enumerate(row):
                        cells[col_index].text = str(value)
                        cells[col_index].width = Inches(float(requested_widths[col_index]))
                        cells[col_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        if row_index < header_rows:
                            for run in cells[col_index].paragraphs[0].runs:
                                run.bold = True; run.font.color.rgb = RGBColor.from_string("17365D")
                            shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "F2F4F7"); cells[col_index]._tc.get_or_add_tcPr().append(shading)
                _set_docx_table_geometry(word_table, requested_widths, total_dxa=9360, indent_dxa=120)
                if isinstance(table_spec, dict) and table_spec.get("caption"):
                    caption = document.add_paragraph(style="Caption")
                    caption.add_run(str(table_spec.get("label") or "Table") + " ").bold = True
                    caption.add_run(str(table_spec["caption"]))
                    if table_spec.get("bookmark"):
                        _bookmark_caption_label(caption, str(table_spec["bookmark"]))
                    caption.paragraph_format.keep_with_next = False
            for formula_index, formula in enumerate(item.get("formulas") or [], 1):
                spec = formula if isinstance(formula, dict) else {"text": formula}
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph._p.append(_omml_equation(str(spec.get("text") or spec.get("formula") or "")))
                if spec.get("label"):
                    paragraph.add_run(f"  ({spec['label']})")
            for figure_index, figure in enumerate(item.get("figures") or [], 1):
                spec = figure if isinstance(figure, dict) else {"path": figure}
                image_path = Path(str(spec.get("path") or ""))
                if image_path.is_file():
                    paragraph = document.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = paragraph.add_run()
                    run.add_picture(str(image_path), width=Inches(float(spec.get("width_inches") or 5.8)))
                    doc_pr = paragraph._p.find(".//wp:docPr", {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"})
                    if doc_pr is not None:
                        doc_pr.set("descr", str(spec.get("alt_text") or spec.get("caption") or image_path.name))
                        doc_pr.set("title", str(spec.get("label") or f"Figure {figure_index}"))
                    if spec.get("caption"):
                        caption = document.add_paragraph(style="Caption")
                        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        caption.add_run(str(spec.get("label") or f"Figure {figure_index}") + " ").bold = True
                        caption.add_run(str(spec["caption"]))
                        if spec.get("bookmark"):
                            _bookmark_caption_label(caption, str(spec["bookmark"]))
        header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_run = header.add_run(str((metadata or {}).get("header") or title or target.stem)); header_run.font.name = "Calibri"; header_run.font.size = Pt(8); header_run.font.color.rgb = RGBColor.from_string("64748B")
        footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run(str((metadata or {}).get("footer") or "ClawOS • Auditable Office Artifact") + "  |  ")
        _append_word_field(footer, "PAGE")
        document.core_properties.title = str(title or target.stem)
        document.core_properties.subject = str((metadata or {}).get("subject") or "ClawOS Office Artifact")
        document.save(target)
        if (metadata or {}).get("footnotes"):
            _patch_docx_footnotes(target, list((metadata or {}).get("footnotes") or []))
        return str(target)
    except ImportError:
        return _create_docx_stdlib(path, title, sections)


def _create_docx_stdlib(path, title, sections):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    paragraphs = [(str(title or target.stem), "Title")]
    for section in sections or []:
        if section.get("heading"):
            paragraphs.append((str(section["heading"]), "Heading1"))
        paragraphs.extend((str(item), "Normal") for item in (section.get("paragraphs") or []))
        paragraphs.extend((f"• {item}", "Normal") for item in (section.get("bullets") or []))
        for row in section.get("table") or []:
            paragraphs.append((" | ".join(str(value) for value in row), "Normal"))
    body = []
    for text, style in paragraphs:
        style_xml = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style != "Normal" else ""
        body.append(f'<w:p>{style_xml}<w:r><w:t xml:space="preserve">{html.escape(text)}</w:t></w:r></w:p>')
    document_xml = _xml_header(
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:body>{"".join(body)}<w:sectPr/></w:body></w:document>'
    )
    content_types = _xml_header(
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '</Types>'
    )
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", _rels({"rId1": ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument", "word/document.xml")}))
        archive.writestr("word/document.xml", document_xml)
    return str(target)


def replace_docx_text(path, replacements, output_path=None):
    from services.office_roundtrip_service import OfficeRoundtripService
    source = Path(path)
    output = Path(output_path) if output_path else source
    OfficeRoundtripService().edit_docx(source, output, replacements=replacements or {})
    return str(output)


def _append_word_field(paragraph, instruction):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText"); code.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); code.text = f" {instruction} "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for node in (begin, code, separate, text, end): run.append(node)
    paragraph._p.append(run)


def _add_docx_paragraph(document, value):
    """Add text while preserving explicit Word REF fields.

    `[[REF:figure_1]]` becomes a real REF field rather than display-only text.
    """
    paragraph = document.add_paragraph()
    cursor = 0
    for match in re.finditer(r"\[\[REF:([A-Za-z_][A-Za-z0-9_.-]*)\]\]", str(value or "")):
        if match.start() > cursor:
            paragraph.add_run(str(value)[cursor:match.start()])
        _append_word_field(paragraph, f"REF {match.group(1)} \\h")
        cursor = match.end()
    if cursor < len(str(value or "")):
        paragraph.add_run(str(value)[cursor:])
    return paragraph


def _bookmark_paragraph(paragraph, name):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", str(name or ""))[:40]
    if not safe or not re.match(r"[A-Za-z_]", safe):
        raise ValueError("bookmark must start with a letter or underscore")
    # Stable per-name ids are sufficient within generated documents and make
    # deterministic OOXML round-trip tests possible.
    bookmark_id = str(int.from_bytes(safe.encode("utf-8"), "little") % 2147483647)
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), bookmark_id); start.set(qn("w:name"), safe)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(0, start); paragraph._p.append(end)


def _bookmark_caption_label(paragraph, name):
    """Bookmark only the caption label run so REF does not duplicate caption prose."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", str(name or ""))[:40]
    if not safe or not re.match(r"[A-Za-z_]", safe):
        raise ValueError("bookmark must start with a letter or underscore")
    bookmark_id = str(int.from_bytes(safe.encode("utf-8"), "little") % 2147483647)
    runs = paragraph._p.findall(qn("w:r"))
    if not runs:
        return _bookmark_paragraph(paragraph, safe)
    label_run = runs[0]; index = list(paragraph._p).index(label_run)
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), bookmark_id); start.set(qn("w:name"), safe)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(index, start); paragraph._p.insert(index + 2, end)


def _omml_equation(value):
    from docx.oxml import OxmlElement
    node = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t"); text.text = str(value or "")
    run.append(text); math.append(run); node.append(math)
    return node


def _set_docx_table_geometry(table, widths_inches, *, total_dxa=9360, indent_dxa=120):
    """Encode fixed DXA table width, indent, grid, cell widths, and margins."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    raw = [max(1.0, float(value)) for value in widths_inches]
    scale = total_dxa / sum(raw)
    widths = [int(round(value * scale)) for value in raw]
    widths[-1] += total_dxa - sum(widths)
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblCellMar"):
        existing = tbl_pr.find(qn(tag))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW"); tbl_w.set(qn("w:w"), str(total_dxa)); tbl_w.set(qn("w:type"), "dxa"); tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd"); tbl_ind.set(qn("w:w"), str(indent_dxa)); tbl_ind.set(qn("w:type"), "dxa"); tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tbl_pr.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        node = OxmlElement(f"w:{side}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); margins.append(node)
    tbl_pr.append(margins)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol"); column.set(qn("w:w"), str(value)); grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW"); tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:w"), str(widths[min(index, len(widths) - 1)])); tc_w.set(qn("w:type"), "dxa")


def _patch_docx_footnotes(path: Path, footnotes: list[dict]):
    parts = _read_zip(path)
    root = ET.fromstring(parts["word/document.xml"])
    # ElementTree does not retain namespace declarations referenced only from
    # mc:Ignorable values. Removing this optional hint avoids producing a
    # Word-rejected package with undeclared w14/wp14 prefixes.
    root.attrib.pop(f"{{{MC_NS}}}Ignorable", None)
    notes_root = ET.fromstring(parts.get("word/footnotes.xml", (
        f'<w:footnotes xmlns:w="{WML_NS}">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '</w:footnotes>'
    ).encode("utf-8")))
    next_id = max([int(item.get(f"{{{WML_NS}}}id", "0")) for item in notes_root.findall(f"{{{WML_NS}}}footnote")] + [0]) + 1
    for item in footnotes:
        marker = str(item.get("marker") or item.get("key") or "").strip()
        marker = marker if marker.startswith("[[FN:") else f"[[FN:{marker}]]"
        matched = False
        for parent in root.iter():
            for index, run in enumerate(list(parent)):
                texts = list(run.iter(f"{{{WML_NS}}}t")) if run.tag == f"{{{WML_NS}}}r" else []
                if len(texts) != 1 or marker not in str(texts[0].text or ""):
                    continue
                before, _, after = str(texts[0].text).partition(marker)
                texts[0].text = before
                if before.startswith(" ") or before.endswith(" "):
                    texts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                reference_run = ET.Element(f"{{{WML_NS}}}r")
                reference_properties = ET.SubElement(reference_run, f"{{{WML_NS}}}rPr")
                ET.SubElement(reference_properties, f"{{{WML_NS}}}vertAlign", {f"{{{WML_NS}}}val": "superscript"})
                ET.SubElement(reference_run, f"{{{WML_NS}}}footnoteReference", {f"{{{WML_NS}}}id": str(next_id)})
                parent.insert(index + 1, reference_run)
                if after:
                    after_run = ET.Element(f"{{{WML_NS}}}r")
                    after_text = ET.SubElement(after_run, f"{{{WML_NS}}}t"); after_text.text = after
                    if after.startswith(" ") or after.endswith(" "):
                        after_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    parent.insert(index + 2, after_run)
                matched = True
                break
            if matched: break
        if not matched:
            raise ValueError(f"footnote marker not found in a single Word run: {marker}")
        note = ET.SubElement(notes_root, f"{{{WML_NS}}}footnote", {f"{{{WML_NS}}}id": str(next_id)})
        paragraph = ET.SubElement(note, f"{{{WML_NS}}}p")
        reference_run = ET.SubElement(paragraph, f"{{{WML_NS}}}r")
        reference_properties = ET.SubElement(reference_run, f"{{{WML_NS}}}rPr")
        ET.SubElement(reference_properties, f"{{{WML_NS}}}vertAlign", {f"{{{WML_NS}}}val": "superscript"})
        ET.SubElement(reference_run, f"{{{WML_NS}}}footnoteRef")
        text_run = ET.SubElement(paragraph, f"{{{WML_NS}}}r")
        ET.SubElement(text_run, f"{{{WML_NS}}}t", {"{http://www.w3.org/XML/1998/namespace}space": "preserve"}).text = " " + str(item.get("text") or "")
        next_id += 1
    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    parts["word/footnotes.xml"] = ET.tostring(notes_root, encoding="utf-8", xml_declaration=True)
    types = ET.fromstring(parts["[Content_Types].xml"])
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    if not any(node.get("PartName") == "/word/footnotes.xml" for node in types):
        ET.SubElement(types, f"{{{ct_ns}}}Override", {"PartName": "/word/footnotes.xml", "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"})
    ET.register_namespace("", ct_ns)
    parts["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
    ET.register_namespace("", MAIN_NS)
    rel_path = "word/_rels/document.xml.rels"
    rels = ET.fromstring(parts[rel_path])
    if not any(node.get("Type", "").endswith("/footnotes") for node in rels):
        used = {node.get("Id") for node in rels}; index = 1
        while f"rId{index}" in used: index += 1
        ET.SubElement(rels, f"{{{PKG_REL_NS}}}Relationship", {"Id": f"rId{index}", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes", "Target": "footnotes.xml"})
    ET.register_namespace("", PKG_REL_NS)
    parts[rel_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    ET.register_namespace("", MAIN_NS)
    _write_zip(path, parts)


def create_pdf(path, title="", sections=None, metadata=None):
    """Create a polished, searchable PDF using the installed ReportLab engine."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError as exc:
        raise RuntimeError("reportlab is required for PDF creation") from exc
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    font_name = "Helvetica"
    document_text = str(title or "") + " " + " ".join(
        str(value)
        for section in (sections or [])
        for value in ([section.get("heading", "")] + list(section.get("paragraphs") or []) + list(section.get("bullets") or []))
    )
    needs_cjk = bool(re.search(r"[\u3400-\u9fff]", document_text))
    font_candidates = (
        ((Path("C:/Windows/Fonts/NotoSansSC-VF.ttf"), "NotoSansSC"), (Path("C:/Windows/Fonts/msyh.ttc"), "MicrosoftYaHei"))
        if needs_cjk else
        ((Path("C:/Windows/Fonts/arial.ttf"), "ArialEmbedded"), (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), "DejaVuSans"))
    )
    for font_path, candidate in font_candidates:
        if font_path.is_file():
            try:
                pdfmetrics.registerFont(TTFont(candidate, str(font_path), subfontIndex=0)); font_name=candidate; break
            except Exception:
                continue
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ArtifactTitle", parent=styles["Title"], fontName=font_name, alignment=TA_CENTER, fontSize=22, leading=27, textColor=colors.HexColor("#172033"), spaceAfter=12 * mm))
    styles.add(ParagraphStyle(name="ArtifactHeading", parent=styles["Heading2"], fontName=font_name, fontSize=15, leading=20, textColor=colors.HexColor("#17365D"), spaceAfter=3 * mm))
    styles.add(ParagraphStyle(name="ArtifactBody", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16, textColor=colors.HexColor("#243047"), spaceAfter=3 * mm))
    styles.add(ParagraphStyle(name="ArtifactTableHeader", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16, textColor=colors.white, spaceAfter=0))
    styles.add(ParagraphStyle(name="ArtifactCaption", parent=styles["BodyText"], fontName=font_name, fontSize=9, leading=13, textColor=colors.HexColor("#475569"), alignment=TA_CENTER, spaceAfter=3 * mm))
    styles.add(ParagraphStyle(name="ArtifactFormula", parent=styles["Code"], fontName="Courier", fontSize=10, leading=14, alignment=TA_CENTER, spaceBefore=2 * mm, spaceAfter=2 * mm))
    doc = SimpleDocTemplate(str(target), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=str(title or target.stem), author="ClawOS")
    story = [Paragraph(html.escape(str(title or target.stem)), styles["ArtifactTitle"])]
    for item in sections or []:
        if item.get("page_break") and len(story) > 1:
            story.append(PageBreak())
        heading = str(item.get("heading") or "").strip()
        if heading:
            story.extend([Paragraph(html.escape(heading), styles["ArtifactHeading"]), Spacer(1, 2 * mm)])
        for paragraph in item.get("paragraphs") or []:
            story.append(Paragraph(html.escape(str(paragraph)).replace("\n", "<br/>"), styles["ArtifactBody"]))
        for bullet in item.get("bullets") or []:
            story.append(Paragraph(f"• {html.escape(str(bullet))}", styles["ArtifactBody"]))
        table_spec = item.get("table") or []
        table = (table_spec.get("rows") or []) if isinstance(table_spec, dict) else table_spec
        if table:
            rows = [[Paragraph(html.escape(str(value)), styles["ArtifactTableHeader"] if row_index == 0 else styles["ArtifactBody"]) for value in row] for row_index, row in enumerate(table)]
            requested_widths = list(table_spec.get("column_widths") or []) if isinstance(table_spec, dict) else []
            col_widths = [float(value) * mm for value in requested_widths] if len(requested_widths) == max(len(row) for row in table) else None
            artifact_table = Table(rows, repeatRows=int(table_spec.get("header_rows", 1)) if isinstance(table_spec, dict) else 1, hAlign="LEFT", colWidths=col_widths)
            artifact_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#172033")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.extend([artifact_table, Spacer(1, 4 * mm)])
            if isinstance(table_spec, dict) and table_spec.get("caption"):
                story.append(Paragraph(html.escape(str(table_spec.get("label") or "Table") + " " + str(table_spec["caption"])), styles["ArtifactCaption"]))
        for formula_index, formula in enumerate(item.get("formulas") or [], 1):
            spec = formula if isinstance(formula, dict) else {"text": formula}
            value = str(spec.get("text") or spec.get("formula") or "")
            label = f"   ({spec['label']})" if spec.get("label") else ""
            story.append(Preformatted(value + label, styles["ArtifactFormula"]))
        for figure_index, figure in enumerate(item.get("figures") or [], 1):
            spec = figure if isinstance(figure, dict) else {"path": figure}
            image_path = Path(str(spec.get("path") or ""))
            if image_path.is_file():
                image = Image(str(image_path))
                max_width, max_height = 165 * mm, 100 * mm
                ratio = min(max_width / max(1, image.imageWidth), max_height / max(1, image.imageHeight), 1.0)
                image.drawWidth = image.imageWidth * ratio; image.drawHeight = image.imageHeight * ratio
                caption = Paragraph(html.escape(str(spec.get("label") or f"Figure {figure_index}") + " " + str(spec.get("caption") or "")), styles["ArtifactCaption"])
                story.append(KeepTogether([image, Spacer(1, 1.5 * mm), caption]))
    footnotes = list((metadata or {}).get("footnotes") or [])
    if footnotes:
        story.extend([Spacer(1, 5 * mm), Paragraph("Notes", styles["ArtifactHeading"])])
        for index, note in enumerate(footnotes, 1):
            story.append(Paragraph(f"<super>{index}</super> {html.escape(str(note.get('text') or ''))}", styles["ArtifactBody"]))

    def draw_footer(canvas, document):
        canvas.saveState(); canvas.setFont(font_name, 8); canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawCentredString(A4[0] / 2, 9 * mm, f"{(metadata or {}).get('footer') or 'ClawOS - Auditable Office Artifact'}  |  {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return str(target)


def _xml_header(text):
    return '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + text


def _col_name(index):
    index = int(index)
    name = ""
    while index:
        index, rem = divmod(index - 1, 26)
        name = chr(65 + rem) + name
    return name


def _cell_ref(row_index, col_index):
    return f"{_col_name(col_index)}{int(row_index)}"


def _cell_kind(value):
    if isinstance(value, bool):
        return "b", "1" if value else "0"
    if isinstance(value, (int, float)):
        return "n", str(value)
    if isinstance(value, str) and value.startswith("="):
        return "f", value[1:]
    return "inlineStr", str(value)


def create_xlsx(path, sheets):
    professional = _create_xlsx_openpyxl(path, sheets)
    if professional:
        return professional
    return _create_xlsx_stdlib(path, sheets)


def _create_xlsx_openpyxl(path, sheets):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except Exception:
        return ""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    sheets = sheets or [{"name": "Sheet1", "rows": []}]
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet in sheets:
        ws = wb.create_sheet(str(sheet.get("name") or "Sheet1")[:31])
        for row_index, row in enumerate(sheet.get("rows", []) or [], 1):
            for col_index, value in enumerate(row, 1):
                cell = ws.cell(row=row_index, column=col_index, value=value)
                if row_index == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F2937")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col_index in range(1, ws.max_column + 1):
            width = max(10, min(36, max(len(str(ws.cell(row=row, column=col_index).value or "")) for row in range(1, ws.max_row + 1)) + 2))
            ws.column_dimensions[get_column_letter(col_index)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        if ws.max_row >= 2 and ws.max_column >= 1:
            table = Table(displayName=f"Table{len(wb.worksheets)}", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
            ws.add_table(table)
        for item in sheet.get("formulas", []) or []:
            ws[str(item.get("cell") or "A1")] = str(item.get("formula") or "")
            ws[str(item.get("cell") or "A1")].number_format = str(item.get("number_format") or "#,##0.00")
        chart_cfg = sheet.get("chart") or {}
        if chart_cfg and ws.max_row > 1:
            chart=BarChart(); chart.title=str(chart_cfg.get("title") or ws.title)
            data_col=max(2,min(int(chart_cfg.get("data_column") or 2),ws.max_column)); category_col=max(1,min(int(chart_cfg.get("category_column") or 1),ws.max_column))
            chart.add_data(Reference(ws,min_col=data_col,min_row=1,max_row=ws.max_row),titles_from_data=True)
            chart.set_categories(Reference(ws,min_col=category_col,min_row=2,max_row=ws.max_row)); ws.add_chart(chart,str(chart_cfg.get("anchor") or "F2"))
    wb.save(path)
    return str(path)


def _create_xlsx_stdlib(path, sheets):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    sheets = sheets or [{"name": "Sheet1", "rows": []}]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _xlsx_content_types(len(sheets)))
        z.writestr("_rels/.rels", _rels({
            "rId1": ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument", "xl/workbook.xml"),
            "rId2": ("http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties", "docProps/core.xml"),
            "rId3": ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties", "docProps/app.xml"),
        }))
        z.writestr("docProps/core.xml", _core_props())
        z.writestr("docProps/app.xml", _app_props())
        z.writestr("xl/workbook.xml", _workbook_xml(sheets))
        z.writestr("xl/_rels/workbook.xml.rels", _workbook_rels(len(sheets)))
        z.writestr("xl/styles.xml", _xlsx_styles())
        for index, sheet in enumerate(sheets, 1):
            z.writestr(f"xl/worksheets/sheet{index}.xml", _worksheet_xml(sheet.get("rows", [])))
    return str(path)


def edit_xlsx(path, updates, output_path=None):
    from services.office_roundtrip_service import OfficeRoundtripService
    source = Path(path)
    output = Path(output_path) if output_path else source
    OfficeRoundtripService().edit_xlsx(source, output, updates or [])
    return str(output)


def _edit_xlsx_openpyxl(path, updates, output_path=None):
    try:
        from openpyxl import load_workbook
    except Exception:
        return ""
    path = Path(path)
    output = Path(output_path) if output_path else path
    wb = load_workbook(path)
    for item in updates or []:
        sheet_key = str(item.get("sheet") or wb.sheetnames[0])
        ws = wb[sheet_key] if sheet_key in wb.sheetnames else wb[wb.sheetnames[0]]
        ws[str(item.get("cell", "A1")).upper()] = item.get("value", "")
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return str(output)


def _edit_xlsx_stdlib(path, updates, output_path=None):
    path = Path(path)
    output = Path(output_path) if output_path else path
    files = _read_zip(path)
    sheet_map = _xlsx_sheet_map(files)
    for item in updates or []:
        sheet_key = str(item.get("sheet") or "Sheet1")
        sheet_file = sheet_map.get(sheet_key) or sheet_map.get("1") or "xl/worksheets/sheet1.xml"
        files[sheet_file] = _update_worksheet(files[sheet_file], item.get("cell", "A1"), item.get("value", ""))
    _write_zip(output, files)
    return str(output)


def create_pptx(path, title="", slides=None):
    professional = _create_pptx_python_pptx(path, title, slides)
    if professional:
        return professional
    return _create_pptx_stdlib(path, title, slides)


def _create_pptx_python_pptx(path, title="", slides=None):
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except Exception:
        return ""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    slides = slides or [{"title": title or "ClawOS Deck", "bullets": []}]
    prs = Presentation()
    for slide_payload in slides:
        layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        background=slide.background.fill; background.solid(); background.fore_color.rgb=RGBColor(248,250,252)
        title_box=slide.shapes.add_textbox(Inches(.65),Inches(.48),Inches(8.7),Inches(.72)); title_frame=title_box.text_frame; title_frame.clear(); p=title_frame.paragraphs[0]; p.text=str(slide_payload.get("title") or title or "ClawOS"); p.font.size=Pt(30); p.font.bold=True; p.font.color.rgb=RGBColor(23,54,93)
        body = slide.shapes.add_textbox(Inches(.8),Inches(1.55),Inches(8.3),Inches(4.7))
        body.text_frame.clear(); body.text_frame.word_wrap=True
        for bullet in slide_payload.get("bullets", []) or []:
            paragraph = body.text_frame.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.level = 0
            paragraph.font.size = Pt(20); paragraph.font.color.rgb=RGBColor(36,48,71); paragraph.space_after=Pt(12)
        if slide_payload.get("notes"):
            box = slide.shapes.add_textbox(Inches(0.7), Inches(6.65), Inches(8.6), Inches(0.25))
            box.text_frame.text = str(slide_payload["notes"])
            box.text_frame.paragraphs[0].font.size=Pt(10); box.text_frame.paragraphs[0].font.color.rgb=RGBColor(100,116,139)
    if len(prs.slides) > len(slides):
        xml_slides = prs.slides._sldIdLst
        rel_id = xml_slides[0].rId
        prs.part.drop_rel(rel_id)
        xml_slides.remove(xml_slides[0])
    prs.save(path)
    return str(path)


def _create_pptx_stdlib(path, title="", slides=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    slides = slides or [{"title": title or "ClawOS Deck", "bullets": []}]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _pptx_content_types(len(slides)))
        z.writestr("_rels/.rels", _rels({
            "rId1": ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument", "ppt/presentation.xml"),
            "rId2": ("http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties", "docProps/core.xml"),
            "rId3": ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties", "docProps/app.xml"),
        }))
        z.writestr("docProps/core.xml", _core_props())
        z.writestr("docProps/app.xml", _app_props())
        z.writestr("ppt/presentation.xml", _presentation_xml(len(slides)))
        z.writestr("ppt/_rels/presentation.xml.rels", _presentation_rels(len(slides)))
        z.writestr("ppt/slideMasters/slideMaster1.xml", _blank_part("p:sldMaster"))
        z.writestr("ppt/slideLayouts/slideLayout1.xml", _blank_part("p:sldLayout"))
        z.writestr("ppt/theme/theme1.xml", _theme_xml())
        for index, slide in enumerate(slides, 1):
            z.writestr(f"ppt/slides/slide{index}.xml", _slide_xml(slide, index))
            z.writestr(f"ppt/slides/_rels/slide{index}.xml.rels", _rels({}))
    return str(path)


def replace_pptx_text(path, replacements, output_path=None):
    from services.office_roundtrip_service import OfficeRoundtripService
    source = Path(path)
    output = Path(output_path) if output_path else source
    OfficeRoundtripService().edit_pptx(source, output, replacements or {})
    return str(output)


def _replace_pptx_text_python_pptx(path, replacements, output_path=None):
    try:
        from pptx import Presentation
    except Exception:
        return ""
    path = Path(path)
    output = Path(output_path) if output_path else path
    prs = Presentation(str(path))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    text = run.text
                    for old, new in (replacements or {}).items():
                        text = text.replace(str(old), str(new))
                    run.text = text
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return str(output)


def _replace_pptx_text_stdlib(path, replacements, output_path=None):
    path = Path(path)
    output = Path(output_path) if output_path else path
    files = _read_zip(path)
    for name, content in list(files.items()):
        if not re.match(r"ppt/slides/slide\d+\.xml$", name):
            continue
        text = content.decode("utf-8")
        for old, new in (replacements or {}).items():
            text = text.replace(html.escape(str(old)), html.escape(str(new)))
            text = text.replace(str(old), html.escape(str(new)))
        files[name] = text.encode("utf-8")
    _write_zip(output, files)
    return str(output)


def _read_zip(path):
    with zipfile.ZipFile(path, "r") as z:
        return {name: z.read(name) for name in z.namelist()}


def _write_zip(path, files):
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in files.items():
            z.writestr(name, content)


def _xlsx_sheet_map(files):
    root = ET.fromstring(files["xl/workbook.xml"])
    mapping = {}
    for index, sheet in enumerate(root.findall(f".//{{{MAIN_NS}}}sheet"), 1):
        name = sheet.attrib.get("name", f"Sheet{index}")
        mapping[name] = f"xl/worksheets/sheet{index}.xml"
        mapping[str(index)] = f"xl/worksheets/sheet{index}.xml"
    return mapping


def _update_worksheet(content, cell, value):
    root = ET.fromstring(content)
    sheet_data = root.find(f"{{{MAIN_NS}}}sheetData")
    if sheet_data is None:
        sheet_data = ET.SubElement(root, f"{{{MAIN_NS}}}sheetData")
    match = re.match(r"([A-Z]+)(\d+)$", str(cell).upper())
    if not match:
        raise ValueError(f"invalid cell reference: {cell}")
    row_number = int(match.group(2))
    row = None
    for candidate in sheet_data.findall(f"{{{MAIN_NS}}}row"):
        if int(candidate.attrib.get("r", 0)) == row_number:
            row = candidate
            break
    if row is None:
        row = ET.SubElement(sheet_data, f"{{{MAIN_NS}}}row", {"r": str(row_number)})
    target = None
    for candidate in row.findall(f"{{{MAIN_NS}}}c"):
        if candidate.attrib.get("r") == str(cell).upper():
            target = candidate
            break
    if target is None:
        target = ET.SubElement(row, f"{{{MAIN_NS}}}c", {"r": str(cell).upper()})
    target.clear()
    target.attrib["r"] = str(cell).upper()
    kind, serial = _cell_kind(value)
    if kind == "inlineStr":
        target.attrib["t"] = "inlineStr"
        inline = ET.SubElement(target, f"{{{MAIN_NS}}}is")
        ET.SubElement(inline, f"{{{MAIN_NS}}}t").text = serial
    elif kind == "f":
        ET.SubElement(target, f"{{{MAIN_NS}}}f").text = serial
        ET.SubElement(target, f"{{{MAIN_NS}}}v")
    else:
        if kind == "b":
            target.attrib["t"] = "b"
        ET.SubElement(target, f"{{{MAIN_NS}}}v").text = serial
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _rels(items):
    rows = [
        f'<Relationship Id="{rid}" Type="{rel_type}" Target="{target}"/>'
        for rid, (rel_type, target) in items.items()
    ]
    return _xml_header(
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(rows) + "</Relationships>"
    )


def _core_props():
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    return _xml_header(
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" '
        'xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        "<dc:creator>ClawOS</dc:creator><cp:lastModifiedBy>ClawOS</cp:lastModifiedBy>"
        f'<dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>'
        f'<dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>'
        "</cp:coreProperties>"
    )


def _app_props():
    return _xml_header(
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        "<Application>ClawOS</Application></Properties>"
    )


def _xlsx_content_types(sheet_count):
    sheet_parts = "".join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for i in range(1, sheet_count + 1)
    )
    return _xml_header(
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
        '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
        + sheet_parts + "</Types>"
    )


def _workbook_xml(sheets):
    rows = "".join(
        f'<sheet name="{html.escape(str(sheet.get("name") or f"Sheet{i}"))}" sheetId="{i}" r:id="rId{i}"/>'
        for i, sheet in enumerate(sheets, 1)
    )
    return _xml_header(
        f'<workbook xmlns="{MAIN_NS}" xmlns:r="{REL_NS}"><sheets>{rows}</sheets></workbook>'
    )


def _workbook_rels(sheet_count):
    items = {
        f"rId{i}": (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
            f"worksheets/sheet{i}.xml",
        )
        for i in range(1, sheet_count + 1)
    }
    items[f"rId{sheet_count + 1}"] = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles",
        "styles.xml",
    )
    return _rels(items)


def _xlsx_styles():
    return _xml_header(
        f'<styleSheet xmlns="{MAIN_NS}"><fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs></styleSheet>'
    )


def _worksheet_xml(rows):
    row_xml = []
    for row_index, row in enumerate(rows or [], 1):
        cells = []
        for col_index, value in enumerate(row or [], 1):
            ref = _cell_ref(row_index, col_index)
            kind, serial = _cell_kind(value)
            if kind == "inlineStr":
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{html.escape(serial)}</t></is></c>')
            elif kind == "b":
                cells.append(f'<c r="{ref}" t="b"><v>{serial}</v></c>')
            elif kind == "f":
                cells.append(f'<c r="{ref}"><f>{html.escape(serial)}</f><v/></c>')
            else:
                cells.append(f'<c r="{ref}"><v>{serial}</v></c>')
        row_xml.append(f'<row r="{row_index}">{"".join(cells)}</row>')
    return _xml_header(f'<worksheet xmlns="{MAIN_NS}"><sheetData>{"".join(row_xml)}</sheetData></worksheet>')


def _pptx_content_types(slide_count):
    slide_parts = "".join(
        f'<Override PartName="/ppt/slides/slide{i}.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        for i in range(1, slide_count + 1)
    )
    return _xml_header(
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
        '<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>'
        '<Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>'
        '<Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
        '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
        + slide_parts + "</Types>"
    )


def _presentation_xml(slide_count):
    ids = "".join(f'<p:sldId id="{255 + i}" r:id="rId{i}"/>' for i in range(1, slide_count + 1))
    return _xml_header(
        f'<p:presentation xmlns:p="{PML_NS}" xmlns:a="{AML_NS}" xmlns:r="{REL_NS}">'
        f"<p:sldIdLst>{ids}</p:sldIdLst><p:sldSz cx=\"12192000\" cy=\"6858000\" type=\"wide\"/>"
        "</p:presentation>"
    )


def _presentation_rels(slide_count):
    items = {
        f"rId{i}": (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
            f"slides/slide{i}.xml",
        )
        for i in range(1, slide_count + 1)
    }
    items[f"rId{slide_count + 1}"] = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster",
        "slideMasters/slideMaster1.xml",
    )
    items[f"rId{slide_count + 2}"] = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme",
        "theme/theme1.xml",
    )
    return _rels(items)


def _blank_part(tag):
    return _xml_header(
        f'<{tag} xmlns:p="{PML_NS}" xmlns:a="{AML_NS}" xmlns:r="{REL_NS}"><p:cSld><p:spTree>'
        '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr/></p:spTree></p:cSld></' + tag + ">"
    )


def _theme_xml():
    return _xml_header(
        f'<a:theme xmlns:a="{AML_NS}" name="ClawOS"><a:themeElements><a:clrScheme name="ClawOS">'
        '<a:dk1><a:srgbClr val="111827"/></a:dk1><a:lt1><a:srgbClr val="FFFFFF"/></a:lt1>'
        '<a:accent1><a:srgbClr val="2563EB"/></a:accent1></a:clrScheme>'
        '<a:fontScheme name="ClawOS"><a:majorFont><a:latin typeface="Arial"/></a:majorFont>'
        '<a:minorFont><a:latin typeface="Arial"/></a:minorFont></a:fontScheme>'
        '<a:fmtScheme name="ClawOS"/></a:themeElements></a:theme>'
    )


def _slide_xml(slide, index):
    title = html.escape(str(slide.get("title") or f"Slide {index}"))
    bullets = [html.escape(str(item)) for item in (slide.get("bullets") or [])]
    notes = html.escape(str(slide.get("notes") or ""))
    bullet_xml = "".join(
        f'<a:p><a:r><a:rPr lang="zh-CN" sz="2200"/><a:t>{item}</a:t></a:r></a:p>'
        for item in bullets
    )
    if notes:
        bullet_xml += f'<a:p><a:r><a:rPr lang="zh-CN" sz="1600"/><a:t>{notes}</a:t></a:r></a:p>'
    return _xml_header(
        f'<p:sld xmlns:p="{PML_NS}" xmlns:a="{AML_NS}" xmlns:r="{REL_NS}"><p:cSld><p:spTree>'
        '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>'
        f'{_text_box(2, "Title", 700000, 450000, 10800000, 900000, title, 3400, True)}'
        f'{_text_box(3, "Body", 900000, 1550000, 10400000, 4300000, bullet_xml, 2200, False)}'
        '</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>'
    )


def _text_box(shape_id, name, x, y, cx, cy, text, font_size, is_plain_text):
    if is_plain_text:
        body = f'<a:p><a:r><a:rPr lang="zh-CN" sz="{font_size}"/><a:t>{text}</a:t></a:r></a:p>'
    else:
        body = text or '<a:p><a:r><a:t></a:t></a:r></a:p>'
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="square"/><a:lstStyle/>{body}</p:txBody></p:sp>'
    )
